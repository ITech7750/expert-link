from __future__ import annotations

import math
import os
import re
import shutil
import subprocess
import textwrap
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.section import WD_SECTION_START, WD_ORIENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor as DocxRGBColor
from docx.text.parfmt import WD_LINE_SPACING
from lxml import html
from markdown import markdown
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches as PptInches
from pptx.util import Pt as PptPt


ROOT = Path(__file__).resolve().parents[1]
TOOLS_DIR = Path(__file__).resolve().parent
OUTPUT_DIR = ROOT / "final-package"
ASSET_DIR = OUTPUT_DIR / "_assets"

TECH_PPTX = Path("/home/itech/Downloads/Распределённый контур инвентаризации(1).pptx")
BUSINESS_PPTX = Path("/home/itech/Downloads/Распределённый_контур_инвентаризации_оптимизировано_1.pptx")

PROJECT_TITLE = "Распределённая платформа инвентаризации, обмена данными и оперативного взаимодействия для удалённых производственных объектов в условиях нестабильной связи"
PROJECT_SHORT = "Распределённая платформа инвентаризации"
PARTICIPANT = "Титарь Игорь Андреевич"
UNIVERSITY = "НИЯУ МИФИ, ИИКС, кафедра №22"
YEAR = "2026"

BLUE = (13, 93, 184)
DARK = (31, 41, 51)
TEAL = (14, 123, 114)
LIGHT_BLUE = (225, 239, 253)
LIGHT_GRAY = (235, 239, 244)
MID_GRAY = (122, 134, 150)
BORDER = (206, 214, 224)
WHITE = (255, 255, 255)

PPT_BLUE = RGBColor(*BLUE)
PPT_DARK = RGBColor(*DARK)
PPT_TEAL = RGBColor(*TEAL)
PPT_LIGHT_BLUE = RGBColor(*LIGHT_BLUE)
PPT_LIGHT_GRAY = RGBColor(*LIGHT_GRAY)
PPT_MID_GRAY = RGBColor(*MID_GRAY)
PPT_BORDER = RGBColor(*BORDER)
PPT_WHITE = RGBColor(*WHITE)

FINAL_FILES = {
    "application": OUTPUT_DIR / "Газпром_Основная_заявка_Распределённая_платформа_инвентаризации.docx",
    "note": OUTPUT_DIR / "Газпром_Пояснительная_записка_Распределённая_платформа_инвентаризации.docx",
    "lean_canvas": OUTPUT_DIR / "Lean_Canvas_Распределённая_платформа_инвентаризации.docx",
    "summary": OUTPUT_DIR / "Газпром_Summary_Распределённая_платформа_инвентаризации.docx",
    "cover": OUTPUT_DIR / "Газпром_Сопроводительное_письмо.docx",
    "appendices": OUTPUT_DIR / "Газпром_Приложения_и_схемы.docx",
    "sources": OUTPUT_DIR / "Газпром_Источники.docx",
    "business_pptx": OUTPUT_DIR / "Газпром_Бизнес_презентация_Распределённая_платформа_инвентаризации.pptx",
    "technical_pptx": OUTPUT_DIR / "Газпром_Техническая_презентация_Распределённая_платформа_инвентаризации.pptx",
    "checklist": OUTPUT_DIR / "Что_проверить_перед_отправкой.docx",
}

DELIVERY_ALIASES = {
    "application": OUTPUT_DIR / "Заявка_на_проект.docx",
    "note": OUTPUT_DIR / "Пояснительная_записка_по_проекту.docx",
}


def ensure_dirs() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    ASSET_DIR.mkdir(parents=True, exist_ok=True)


def run(cmd: list[str]) -> None:
    subprocess.run(cmd, check=True)


def font_path(bold: bool = False) -> str:
    candidates = [
        "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf" if bold else "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "/usr/share/fonts/truetype/liberation2/LiberationSans-Bold.ttf" if bold else "/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf",
    ]
    for candidate in candidates:
        if Path(candidate).exists():
            return candidate
    raise FileNotFoundError("Не найден системный шрифт для генерации изображений")


FONT_REGULAR = font_path(False)
FONT_BOLD = font_path(True)


def load_font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    return ImageFont.truetype(FONT_BOLD if bold else FONT_REGULAR, size=size)


def wrap_text(draw: ImageDraw.ImageDraw, text: str, font: ImageFont.FreeTypeFont, max_width: int) -> list[str]:
    words = text.split()
    lines: list[str] = []
    current = ""
    for word in words:
        candidate = word if not current else f"{current} {word}"
        if draw.textbbox((0, 0), candidate, font=font)[2] <= max_width:
            current = candidate
        else:
            if current:
                lines.append(current)
            current = word
    if current:
        lines.append(current)
    return lines


def draw_wrapped_text(
    draw: ImageDraw.ImageDraw,
    box: tuple[int, int, int, int],
    text: str,
    font: ImageFont.FreeTypeFont,
    fill: tuple[int, int, int] = DARK,
    align: str = "left",
    spacing: int = 8,
) -> None:
    x0, y0, x1, y1 = box
    max_width = x1 - x0
    lines = wrap_text(draw, text, font, max_width)
    line_heights = [draw.textbbox((0, 0), line, font=font)[3] for line in lines]
    total_h = sum(line_heights) + spacing * max(0, len(lines) - 1)
    cursor_y = y0 + max(0, ((y1 - y0) - total_h) // 2)
    for line, line_h in zip(lines, line_heights):
        bbox = draw.textbbox((0, 0), line, font=font)
        text_w = bbox[2] - bbox[0]
        if align == "center":
            cursor_x = x0 + ((x1 - x0) - text_w) // 2
        elif align == "right":
            cursor_x = x1 - text_w
        else:
            cursor_x = x0
        draw.text((cursor_x, cursor_y), line, font=font, fill=fill)
        cursor_y += line_h + spacing


def draw_round_box(
    draw: ImageDraw.ImageDraw,
    box: tuple[int, int, int, int],
    text: str,
    fill: tuple[int, int, int] = WHITE,
    outline: tuple[int, int, int] = BORDER,
    title_fill: tuple[int, int, int] = DARK,
    title_size: int = 34,
    radius: int = 28,
) -> None:
    draw.rounded_rectangle(box, radius=radius, fill=fill, outline=outline, width=4)
    font = load_font(title_size, bold=True)
    draw_wrapped_text(draw, (box[0] + 20, box[1] + 10, box[2] - 20, box[3] - 10), text, font, fill=title_fill, align="center")


def draw_arrow(draw: ImageDraw.ImageDraw, start: tuple[int, int], end: tuple[int, int], fill: tuple[int, int, int], width: int = 6) -> None:
    draw.line([start, end], fill=fill, width=width)
    angle = math.atan2(end[1] - start[1], end[0] - start[0])
    arrow_len = 18
    left = (
        end[0] - arrow_len * math.cos(angle - math.pi / 6),
        end[1] - arrow_len * math.sin(angle - math.pi / 6),
    )
    right = (
        end[0] - arrow_len * math.cos(angle + math.pi / 6),
        end[1] - arrow_len * math.sin(angle + math.pi / 6),
    )
    draw.polygon([end, left, right], fill=fill)


def diagram_canvas(title: str, subtitle: str | None = None) -> tuple[Image.Image, ImageDraw.ImageDraw]:
    img = Image.new("RGB", (1600, 900), WHITE)
    draw = ImageDraw.Draw(img)
    draw.ellipse((-120, -120, 280, 280), outline=LIGHT_GRAY, width=8)
    draw.ellipse((1320, -120, 1680, 240), outline=LIGHT_GRAY, width=8)
    draw.ellipse((1220, 690, 1680, 1150), outline=LIGHT_GRAY, width=8)
    draw.text((70, 40), title, font=load_font(44, bold=True), fill=BLUE)
    if subtitle:
        draw.text((70, 100), subtitle, font=load_font(24), fill=MID_GRAY)
    return img, draw


def save_image(img: Image.Image, name: str) -> Path:
    path = ASSET_DIR / name
    img.save(path)
    return path


def build_custom_diagrams() -> dict[str, Path]:
    diagrams: dict[str, Path] = {}

    img, draw = diagram_canvas("Разрозненный текущий процесс", "Учёт, обсуждения и материалы живут в разных инструментах")
    draw_round_box(draw, (120, 250, 430, 380), "Центральная\nучётная система", fill=LIGHT_BLUE, outline=BLUE, title_fill=BLUE)
    draw_round_box(draw, (520, 180, 810, 290), "Excel и локальные таблицы", fill=(250, 250, 250))
    draw_round_box(draw, (520, 340, 810, 450), "Бумажные формы и акты", fill=(250, 250, 250))
    draw_round_box(draw, (520, 500, 810, 610), "Сторонний мессенджер", fill=(250, 250, 250))
    draw_round_box(draw, (980, 300, 1440, 470), "Комиссия на удалённой площадке", fill=(233, 247, 245), outline=TEAL, title_fill=TEAL)
    draw_arrow(draw, (430, 315), (520, 235), BLUE)
    draw_arrow(draw, (430, 315), (520, 395), BLUE)
    draw_arrow(draw, (430, 315), (520, 555), BLUE)
    draw_arrow(draw, (810, 235), (980, 350), MID_GRAY)
    draw_arrow(draw, (810, 395), (980, 385), MID_GRAY)
    draw_arrow(draw, (810, 555), (980, 420), MID_GRAY)
    draw_wrapped_text(draw, (980, 520, 1460, 700), "Итог: данные, материалы и обсуждения расходятся, а при потере связи часть процесса уходит в ручной режим.", load_font(28), fill=DARK)
    diagrams["fragmentation"] = save_image(img, "diagram_fragmentation.png")

    img, draw = diagram_canvas("Единый рабочий контур проекта", "Инвентаризация и взаимодействие сотрудников объединены в одной платформе")
    draw_round_box(draw, (550, 220, 1060, 620), "Распределённая платформа", fill=LIGHT_BLUE, outline=BLUE, title_fill=BLUE, title_size=42)
    for box, text in [
        ((650, 300, 960, 380), "Карточки имущества"),
        ((650, 400, 960, 480), "Сессии и проверки"),
        ((110, 280, 420, 390), "Локальные узлы\nи автономная работа"),
        ((110, 470, 420, 580), "Маркировка,\nQR и штрихкоды"),
        ((1190, 280, 1490, 390), "Чаты, файлы,\nобсуждения и звонки"),
        ((1190, 470, 1490, 580), "Central sync\nпри доступной связи"),
    ]:
        draw_round_box(draw, box, text, fill=WHITE)
    draw_arrow(draw, (420, 335), (550, 335), TEAL)
    draw_arrow(draw, (420, 525), (550, 525), TEAL)
    draw_arrow(draw, (1060, 335), (1190, 335), TEAL)
    draw_arrow(draw, (1060, 525), (1190, 525), TEAL)
    diagrams["single_contour"] = save_image(img, "diagram_single_contour.png")

    img, draw = diagram_canvas("Режимы работы с интернетом и без него", "Ключевая идея: отсутствие внешнего канала не останавливает рабочий процесс")
    draw_round_box(draw, (90, 180, 760, 660), "Связь с центральным контуром доступна", fill=LIGHT_BLUE, outline=BLUE, title_fill=BLUE, title_size=38)
    draw_round_box(draw, (840, 180, 1510, 660), "Внешняя связь недоступна", fill=(233, 247, 245), outline=TEAL, title_fill=TEAL, title_size=38)
    left_items = [
        "локальная работа пользователя",
        "доступ к карточкам и сессиям",
        "обсуждения, файлы и звонки",
        "синхронизация с центральным контуром",
        "консолидация и обновление данных",
    ]
    right_items = [
        "локальная автономная работа",
        "взаимодействие между соседними узлами",
        "фиксация результатов и материалов на месте",
        "накопление изменений в очереди",
        "поздняя синхронизация после восстановления связи",
    ]
    for idx, item in enumerate(left_items):
        draw_round_box(draw, (140, 280 + idx * 65, 700, 335 + idx * 65), item, fill=WHITE, radius=18, title_size=24)
    for idx, item in enumerate(right_items):
        draw_round_box(draw, (890, 280 + idx * 65, 1450, 335 + idx * 65), item, fill=WHITE, radius=18, title_size=24)
    diagrams["online_offline"] = save_image(img, "diagram_online_offline.png")

    img, draw = diagram_canvas("Сценарий работы комиссии", "От подготовки сессии до последующей синхронизации")
    steps = [
        "Подготовка\nсессии",
        "Выезд\nна объект",
        "Сканирование\nи открытие карточки",
        "Проверка состояния\nи фиксация расхождений",
        "Обсуждение,\nфото и документы",
        "Локальное сохранение\nи синхронизация",
    ]
    x = 90
    for idx, step in enumerate(steps):
        draw_round_box(draw, (x, 330, x + 220, 510), step, fill=LIGHT_BLUE if idx in (0, 5) else WHITE, outline=BLUE if idx in (0, 5) else BORDER, title_fill=BLUE if idx in (0, 5) else DARK, title_size=26)
        if idx < len(steps) - 1:
            draw_arrow(draw, (x + 220, 420), (x + 260, 420), TEAL)
        x += 250
    diagrams["commission"] = save_image(img, "diagram_commission_workflow.png")

    img, draw = diagram_canvas("Синхронизация изменений", "Локальная очередь и последующая консолидация в центральном контуре")
    participants = [
        ("Пользователь", 160),
        ("Локальный узел", 520),
        ("Очередь изменений", 900),
        ("Центральный контур", 1280),
    ]
    for title, center_x in participants:
        draw_round_box(draw, (center_x - 120, 150, center_x + 120, 230), title, fill=LIGHT_BLUE, outline=BLUE, title_fill=BLUE, title_size=24, radius=18)
        draw.line((center_x, 230, center_x, 720), fill=BORDER, width=3)
    arrows = [
        ((160, 300), (520, 300), "создание / изменение объекта"),
        ((520, 380), (900, 380), "фиксация локального события"),
        ((900, 470), (1280, 470), "upload pending changes"),
        ((1280, 560), (900, 560), "ответ: подтверждение / конфликты"),
        ((900, 640), (520, 640), "обновление локального состояния"),
    ]
    for start, end, label in arrows:
        draw_arrow(draw, start, end, TEAL if start[0] < end[0] else BLUE)
        mid_x = (start[0] + end[0]) // 2 - 120
        draw_round_box(draw, (mid_x, start[1] - 28, mid_x + 240, start[1] + 28), label, fill=WHITE, outline=BORDER, title_size=18, radius=12)
    diagrams["sync"] = save_image(img, "diagram_sync_flow.png")

    img, draw = diagram_canvas(
        "Центральный контур: сервисы и взаимодействие",
        "Маршрутизация через gateway-service, профильная конфигурация и единый слой хранения",
    )
    draw_round_box(
        draw,
        (80, 270, 390, 390),
        "Клиентские узлы\nexpert-link\n(Desktop/Android)",
        fill=WHITE,
        outline=BORDER,
        title_size=26,
    )
    draw_round_box(
        draw,
        (460, 270, 690, 390),
        "gateway-service",
        fill=LIGHT_BLUE,
        outline=BLUE,
        title_fill=BLUE,
        title_size=30,
    )
    draw_arrow(draw, (390, 330), (460, 330), TEAL)

    service_boxes = [
        ("organization-\naccess-service", (760, 205, 995, 305)),
        ("inventory-service", (1020, 205, 1255, 305)),
        ("sync-service", (1280, 205, 1515, 305)),
        ("relay-service", (760, 345, 995, 445)),
        ("attachment-service", (1020, 345, 1255, 445)),
        ("export-service", (1280, 345, 1515, 445)),
    ]
    for text, box in service_boxes:
        draw_round_box(draw, box, text, fill=WHITE, outline=BORDER, title_size=22, radius=18)
        draw_arrow(draw, (690, 330), (box[0], (box[1] + box[3]) // 2), BLUE)

    draw_round_box(draw, (760, 80, 995, 160), "config-server", fill=LIGHT_BLUE, outline=BLUE, title_fill=BLUE, title_size=22, radius=18)
    draw_round_box(draw, (1020, 80, 1255, 160), "discovery-service", fill=LIGHT_BLUE, outline=BLUE, title_fill=BLUE, title_size=22, radius=18)
    draw_arrow(draw, (878, 160), (878, 205), TEAL)
    draw_arrow(draw, (1138, 160), (1138, 205), TEAL)
    draw_round_box(draw, (1280, 80, 1515, 160), "config-repo", fill=WHITE, outline=BORDER, title_size=22, radius=18)
    draw_arrow(draw, (1280, 120), (995, 120), MID_GRAY)

    draw_round_box(draw, (760, 560, 1010, 660), "PostgreSQL\nканонические\nданные", fill=LIGHT_BLUE, outline=BLUE, title_fill=BLUE, title_size=22, radius=18)
    draw_round_box(draw, (1035, 560, 1235, 660), "Redis\npresence, TTL,\nrate limit", fill=LIGHT_BLUE, outline=BLUE, title_fill=BLUE, title_size=20, radius=18)
    draw_round_box(draw, (1260, 560, 1515, 660), "Filesystem / S3\nвложения и\nартефакты", fill=LIGHT_BLUE, outline=BLUE, title_fill=BLUE, title_size=21, radius=18)

    draw_arrow(draw, (878, 445), (878, 560), BLUE)
    draw_arrow(draw, (1138, 305), (900, 560), BLUE)
    draw_arrow(draw, (1398, 305), (900, 560), BLUE)
    draw_arrow(draw, (878, 445), (1135, 560), TEAL)
    draw_arrow(draw, (575, 390), (1135, 560), TEAL)
    draw_arrow(draw, (1138, 445), (1388, 560), BLUE)
    draw_arrow(draw, (1398, 445), (1388, 560), BLUE)

    draw_wrapped_text(
        draw,
        (80, 705, 1515, 860),
        "Поток взаимодействия: клиентские узлы обращаются в gateway-service. "
        "Gateway маршрутизирует запросы в профильные микросервисы центрального контура. "
        "Сервисы работают с каноническим состоянием в PostgreSQL, используют Redis для присутствия и служебных сценариев, "
        "а attachment-service и export-service дополнительно используют профильное двоичное хранилище.",
        load_font(22),
        fill=DARK,
    )
    diagrams["central_contour_services"] = save_image(img, "diagram_central_contour_services.png")
    return diagrams


def render_existing_slide_images() -> dict[str, Path]:
    pdf_path = ASSET_DIR / "technical_source.pdf"
    if pdf_path.exists():
        pdf_path.unlink()
    run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(ASSET_DIR), str(TECH_PPTX)])
    generated_pdf = ASSET_DIR / f"{TECH_PPTX.stem}.pdf"
    if generated_pdf.exists() and generated_pdf != pdf_path:
        generated_pdf.rename(pdf_path)

    images: dict[str, Path] = {}
    slide_map = {
        "architecture_existing": 9,
        "inventory_er_existing": 10,
        "comm_er_existing": 11,
        "packet_sequence_existing": 13,
        "call_flow_existing": 17,
        "file_flow_existing": 19,
        "proto_main": 24,
        "proto_quick_actions": 25,
        "proto_inventory_table": 26,
        "proto_inventory_cards": 27,
        "proto_inventory_filters": 28,
        "proto_inventory_edit": 29,
    }
    for name, page in slide_map.items():
        prefix = ASSET_DIR / name
        run([
            "pdftoppm",
            "-r",
            "220",
            "-singlefile",
            "-f",
            str(page),
            "-l",
            str(page),
            "-png",
            str(pdf_path),
            str(prefix),
        ])
        images[name] = prefix.with_suffix(".png")
    return images


def copy_static_images() -> dict[str, Path]:
    mapping = {
        "mobile_home": Path("/home/itech/IdeaProjects/expert-link/docs/img.png"),
        "mobile_profile": Path("/home/itech/IdeaProjects/expert-link/docs/img_2.png"),
        "mobile_transfers": Path("/home/itech/IdeaProjects/expert-link/docs/img_3.png"),
        "mobile_chats": Path("/home/itech/IdeaProjects/expert-link/docs/img_4.png"),
        "mobile_pairing_scan": Path("/home/itech/IdeaProjects/expert-link/docs/10.jpg"),
        "mobile_call": Path("/home/itech/IdeaProjects/expert-link/docs/9.jpg"),
        "desktop_profile": Path("/home/itech/IdeaProjects/expert-link/docs/11.jpg"),
    }
    copied: dict[str, Path] = {}
    for key, src in mapping.items():
        dst = ASSET_DIR / f"{key}{src.suffix.lower()}"
        shutil.copy2(src, dst)
        copied[key] = dst
    return copied


def qn_set_font(style, name: str) -> None:
    style.font.name = name
    style._element.rPr.rFonts.set(qn("w:eastAsia"), name)


def apply_run_font(run, size: int = 12, bold: bool | None = None, italic: bool | None = None) -> None:
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_paragraph_format(
    paragraph,
    *,
    align=WD_ALIGN_PARAGRAPH.JUSTIFY,
    first_indent_cm: float | None = 1.25,
    left_indent_cm: float | None = None,
    line_spacing: float = 1.5,
    space_before_pt: float = 0,
    space_after_pt: float = 0,
    keep_with_next: bool | None = None,
    page_break_before: bool | None = None,
) -> None:
    pf = paragraph.paragraph_format
    paragraph.alignment = align
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = line_spacing
    pf.space_before = Pt(space_before_pt)
    pf.space_after = Pt(space_after_pt)
    pf.first_line_indent = None if first_indent_cm is None else Cm(first_indent_cm)
    if left_indent_cm is not None:
        pf.left_indent = Cm(left_indent_cm)
    if keep_with_next is not None:
        pf.keep_with_next = keep_with_next
    if page_break_before is not None:
        pf.page_break_before = page_break_before


def normalize_caption(text: str) -> str:
    return re.sub(r"^(Рисунок\s+[А-ЯA-Z]?\d+(?:\.\d+)?)\.\s+", r"\1 – ", text)


def normalize_table_title(text: str) -> str:
    return re.sub(r"^(Таблица\s+[А-ЯA-Z]?\d+(?:\.\d+)?)\.\s+", r"\1 – ", text)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def setup_doc_styles(doc: Document) -> None:
    sec = doc.sections[0]
    sec.top_margin = Cm(2)
    sec.bottom_margin = Cm(2)
    sec.left_margin = Cm(3)
    sec.right_margin = Cm(1.5)
    sec.page_width = Cm(21)
    sec.page_height = Cm(29.7)
    sec.header_distance = Cm(1.25)
    sec.footer_distance = Cm(1.25)
    sec.different_first_page_header_footer = True

    normal = doc.styles["Normal"]
    qn_set_font(normal, "Times New Roman")
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.first_line_indent = Cm(1.25)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    for style_name, size in [("Title", 16), ("Heading 1", 14), ("Heading 2", 14), ("Heading 3", 13)]:
        style = doc.styles[style_name]
        qn_set_font(style, "Times New Roman")
        style.font.size = Pt(size)
        style.font.bold = True
        style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
        style.paragraph_format.line_spacing = 1.5
        style.paragraph_format.space_before = Pt(12 if style_name != "Title" else 0)
        style.paragraph_format.space_after = Pt(0)
        style.paragraph_format.first_line_indent = Cm(0)

    doc.styles["Title"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.styles["Heading 1"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.styles["Heading 2"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    doc.styles["Heading 3"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

    for style_name in ["List Bullet", "List Number"]:
        if style_name in doc.styles:
            style = doc.styles[style_name]
            qn_set_font(style, "Times New Roman")
            style.font.size = Pt(12)
            style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
            style.paragraph_format.line_spacing = 1.5
            style.paragraph_format.space_before = Pt(0)
            style.paragraph_format.space_after = Pt(0)
            style.paragraph_format.left_indent = Cm(1.25)
            style.paragraph_format.first_line_indent = Cm(-0.63)


def add_page_number(paragraph) -> None:
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_begin, instr, fld_end])


def setup_footer(doc: Document) -> None:
    for section in doc.sections:
        footer = section.footer
        p = footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        add_page_number(p)


def add_title_page(doc: Document, title: str, subtitle_lines: Iterable[str] | None = None) -> None:
    p = doc.add_paragraph()
    set_paragraph_format(p, align=WD_ALIGN_PARAGRAPH.CENTER, first_indent_cm=None, line_spacing=1.5)
    run = p.add_run("Студенческая олимпиада «Газпром»\nПрофиль «Информационные системы и технологии»")
    apply_run_font(run, 12, bold=True)

    doc.add_paragraph()
    doc.add_paragraph()

    p = doc.add_paragraph()
    set_paragraph_format(p, align=WD_ALIGN_PARAGRAPH.CENTER, first_indent_cm=None, line_spacing=1.5)
    run = p.add_run(title)
    apply_run_font(run, 16, bold=True)

    if subtitle_lines:
        doc.add_paragraph()
        subtitle_lines = list(subtitle_lines)
        for idx, line in enumerate(subtitle_lines):
            sp = doc.add_paragraph()
            set_paragraph_format(sp, align=WD_ALIGN_PARAGRAPH.CENTER, first_indent_cm=None, line_spacing=1.5)
            r = sp.add_run(line)
            apply_run_font(r, 14 if idx == 0 else 12, bold=(idx == 0))

    for _ in range(8):
        doc.add_paragraph()

    bottom = doc.add_paragraph()
    set_paragraph_format(bottom, align=WD_ALIGN_PARAGRAPH.CENTER, first_indent_cm=None, line_spacing=1.5)
    r = bottom.add_run(YEAR)
    apply_run_font(r, 12)

    doc.add_page_break()


def manual_toc(doc: Document, title: str, items: list[tuple[int, str]]) -> None:
    p = doc.add_paragraph(title, style="Heading 1")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for level, item in items:
        para = doc.add_paragraph()
        para.add_run(item)
        set_paragraph_format(
            para,
            align=WD_ALIGN_PARAGRAPH.LEFT,
            first_indent_cm=None,
            left_indent_cm=0 if level == 2 else (0.75 if level == 3 else 1.5),
            line_spacing=1.5,
        )
    doc.add_page_break()


def clean_text(text: str) -> str:
    normalized = text.replace("\xa0", " ")
    normalized = re.sub(r"\s*;\s*", ", ", normalized)
    return re.sub(r"\s+", " ", normalized).strip()


def render_markdown(doc: Document, md_text: str, skip_first_h1: bool = True, heading_injections: dict[str, list[tuple[Path, str, float]]] | None = None) -> None:
    html_text = markdown(md_text, extensions=["tables", "fenced_code"])
    root = html.fragment_fromstring(html_text, create_parent="div")
    skipped_h1 = False
    for child in root:
        tag = child.tag.lower()
        if tag == "h1" and skip_first_h1 and not skipped_h1:
            skipped_h1 = True
            continue
        if tag in {"h1", "h2", "h3", "h4"}:
            level = min(int(tag[1]), 3)
            heading = clean_text(child.text_content())
            p = doc.add_paragraph(heading, style=f"Heading {level}")
            if level > 1:
                set_paragraph_format(p, align=WD_ALIGN_PARAGRAPH.LEFT, first_indent_cm=None, line_spacing=1.5, keep_with_next=True)
            if heading_injections:
                for key, images in heading_injections.items():
                    if key in heading:
                        for img_path, caption, width_cm in images:
                            insert_figure(doc, img_path, caption, width_cm)
        elif tag == "p":
            text = clean_text(child.text_content())
            if text:
                if text.startswith("Таблица "):
                    p = doc.add_paragraph(normalize_table_title(text))
                    set_paragraph_format(p, align=WD_ALIGN_PARAGRAPH.LEFT, first_indent_cm=None, line_spacing=1.0, keep_with_next=True, space_before_pt=6, space_after_pt=3)
                else:
                    p = doc.add_paragraph(text)
                    set_paragraph_format(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, first_indent_cm=1.25, line_spacing=1.5)
        elif tag in {"ul", "ol"}:
            render_list(doc, child, numbered=(tag == "ol"))
        elif tag == "table":
            render_table(doc, child)
        elif tag == "pre":
            continue


def render_list(doc: Document, node, numbered: bool = False, level: int = 0) -> None:
    style = "List Number" if numbered else "List Bullet"
    for li in node.findall("./li"):
        texts = [clean_text(t) for t in li.xpath("./text()") if clean_text(t)]
        text = " ".join(texts) if texts else clean_text(li.text_content())
        text = text.rstrip(" ,")
        if text:
            p = doc.add_paragraph(style=style)
            p.add_run(text)
            set_paragraph_format(
                p,
                align=WD_ALIGN_PARAGRAPH.JUSTIFY,
                first_indent_cm=-0.63,
                left_indent_cm=1.25 + level * 0.63,
                line_spacing=1.5,
            )
        for nested in li:
            if nested.tag.lower() in {"ul", "ol"}:
                render_list(doc, nested, numbered=(nested.tag.lower() == "ol"), level=level + 1)


def render_table(doc: Document, node) -> None:
    rows = node.findall(".//tr")
    if not rows:
        return
    cols_count = max(len(r.findall("./th")) + len(r.findall("./td")) for r in rows)
    table = doc.add_table(rows=len(rows), cols=cols_count)
    table.style = "Table Grid"
    for r_idx, row in enumerate(rows):
        cells = row.findall("./th") + row.findall("./td")
        for c_idx, cell in enumerate(cells):
            target = table.cell(r_idx, c_idx)
            target.text = clean_text(cell.text_content())
            for p in target.paragraphs:
                set_paragraph_format(
                    p,
                    align=WD_ALIGN_PARAGRAPH.CENTER if r_idx == 0 else WD_ALIGN_PARAGRAPH.LEFT,
                    first_indent_cm=None,
                    line_spacing=1.0,
                )
                for run in p.runs:
                    apply_run_font(run, 12, bold=(r_idx == 0))
    doc.add_paragraph()


def insert_figure(doc: Document, image_path: Path, caption: str, width_cm: float = 16.5) -> None:
    p = doc.add_paragraph()
    set_paragraph_format(p, align=WD_ALIGN_PARAGRAPH.CENTER, first_indent_cm=None, line_spacing=1.0, space_before_pt=6, space_after_pt=3)
    run = p.add_run()
    run.add_picture(str(image_path), width=Cm(width_cm))
    cp = doc.add_paragraph()
    set_paragraph_format(cp, align=WD_ALIGN_PARAGRAPH.CENTER, first_indent_cm=None, line_spacing=1.0, space_before_pt=0, space_after_pt=6)
    cr = cp.add_run(normalize_caption(caption))
    apply_run_font(cr, 12, italic=False)


def headings_from_md(md_path: Path) -> list[tuple[int, str]]:
    items: list[tuple[int, str]] = []
    for line in md_path.read_text(encoding="utf-8").splitlines():
        for prefix, level in [("## ", 2), ("### ", 3), ("#### ", 4)]:
            if line.startswith(prefix):
                items.append((level, line[len(prefix):].strip()))
                break
    return items


def postprocess_document(doc: Document, *, chapter_page_breaks: bool = False) -> None:
    first_main_heading_seen = False
    toc_mode = False
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue

        style_name = paragraph.style.name if paragraph.style is not None else ""

        if text == "Содержание":
            toc_mode = True

        if toc_mode and style_name == "Normal":
            if re.match(r"^\d+\.\s", text):
                indent = 0
            elif re.match(r"^\d+\.\d+\s", text):
                indent = 0.75
            else:
                indent = 1.5
            set_paragraph_format(
                paragraph,
                align=WD_ALIGN_PARAGRAPH.LEFT,
                first_indent_cm=None,
                left_indent_cm=indent,
                line_spacing=1.5,
            )
            for run in paragraph.runs:
                apply_run_font(run, 12)
            continue

        if style_name == "Normal" and paragraph.alignment not in {WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.RIGHT}:
            if text.startswith("Таблица "):
                paragraph.text = normalize_table_title(text)
                set_paragraph_format(paragraph, align=WD_ALIGN_PARAGRAPH.LEFT, first_indent_cm=None, line_spacing=1.0, keep_with_next=True, space_before_pt=6, space_after_pt=3)
                for run in paragraph.runs:
                    apply_run_font(run, 12)
            elif not text.startswith("Рисунок "):
                set_paragraph_format(paragraph, align=WD_ALIGN_PARAGRAPH.JUSTIFY, first_indent_cm=1.25, line_spacing=1.5)
                for run in paragraph.runs:
                    apply_run_font(run, 12)

        if style_name == "Heading 1":
            set_paragraph_format(paragraph, align=WD_ALIGN_PARAGRAPH.CENTER, first_indent_cm=None, line_spacing=1.5, space_before_pt=12, keep_with_next=True)
            for run in paragraph.runs:
                apply_run_font(run, 14, bold=True)

        if style_name == "Heading 2":
            is_main_heading = re.match(r"^\d+\.\s", text) is not None
            set_paragraph_format(
                paragraph,
                align=WD_ALIGN_PARAGRAPH.LEFT,
                first_indent_cm=None,
                line_spacing=1.5,
                space_before_pt=12,
                keep_with_next=True,
                page_break_before=(chapter_page_breaks and first_main_heading_seen and is_main_heading),
            )
            if is_main_heading:
                first_main_heading_seen = True
                toc_mode = False
            for run in paragraph.runs:
                apply_run_font(run, 14 if is_main_heading else 13, bold=True)

        if style_name == "Heading 3":
            set_paragraph_format(paragraph, align=WD_ALIGN_PARAGRAPH.LEFT, first_indent_cm=None, line_spacing=1.5, space_before_pt=12, keep_with_next=True)
            for run in paragraph.runs:
                apply_run_font(run, 13, bold=True)


def build_doc_from_md(
    md_name: str,
    out_path: Path,
    title: str,
    subtitle_lines: Iterable[str] | None = None,
    title_page: bool = True,
    manual_contents: bool = False,
    heading_injections: dict[str, list[tuple[Path, str, float]]] | None = None,
    chapter_page_breaks: bool = False,
) -> None:
    md_path = ROOT / md_name
    md_text = md_path.read_text(encoding="utf-8")
    doc = Document()
    setup_doc_styles(doc)
    setup_footer(doc)
    if title_page:
        add_title_page(doc, title, subtitle_lines)
    if manual_contents:
        manual_toc(doc, "Содержание", headings_from_md(md_path))
    render_markdown(doc, md_text, skip_first_h1=True, heading_injections=heading_injections)
    postprocess_document(doc, chapter_page_breaks=chapter_page_breaks)
    doc.save(out_path)


def build_summary_doc() -> None:
    doc = Document()
    setup_doc_styles(doc)
    setup_footer(doc)
    p = doc.add_paragraph(PROJECT_TITLE, style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph()
    render_markdown(doc, (ROOT / "06-summary.md").read_text(encoding="utf-8"), skip_first_h1=True)
    doc.save(FINAL_FILES["summary"])


def add_canvas_block(cell, title: str, body: list[str], *, subtitle: str | None = None, fill: str = "F7FAFC", title_fill: str = "0D5DB8") -> None:
    cell.text = ""
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    set_cell_shading(cell, fill)
    set_cell_margins(cell, top=120, start=140, bottom=120, end=140)

    p = cell.paragraphs[0]
    set_paragraph_format(p, align=WD_ALIGN_PARAGRAPH.LEFT, first_indent_cm=None, line_spacing=1.15, space_before_pt=0, space_after_pt=3, keep_with_next=True)
    r = p.add_run(title)
    apply_run_font(r, 11, bold=True)
    r.font.color.rgb = DocxRGBColor(13, 93, 184)

    if subtitle:
        p2 = cell.add_paragraph()
        set_paragraph_format(p2, align=WD_ALIGN_PARAGRAPH.LEFT, first_indent_cm=None, line_spacing=1.0, space_before_pt=0, space_after_pt=4)
        r2 = p2.add_run(subtitle)
        apply_run_font(r2, 9, italic=True)

    for idx, item in enumerate(body):
        p_item = cell.add_paragraph()
        set_paragraph_format(
            p_item,
            align=WD_ALIGN_PARAGRAPH.LEFT,
            first_indent_cm=-0.45,
            left_indent_cm=0.45,
            line_spacing=1.1,
            space_before_pt=(1 if idx else 2),
            space_after_pt=0,
        )
        run = p_item.add_run(f"• {item}")
        apply_run_font(run, 10)


def build_lean_canvas_doc() -> None:
    doc = Document()
    setup_doc_styles(doc)
    setup_footer(doc)

    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = sec.page_height, sec.page_width
    sec.left_margin = Cm(1.5)
    sec.right_margin = Cm(1.5)
    sec.top_margin = Cm(1.5)
    sec.bottom_margin = Cm(1.5)

    title = doc.add_paragraph()
    set_paragraph_format(title, align=WD_ALIGN_PARAGRAPH.CENTER, first_indent_cm=None, line_spacing=1.15, space_before_pt=0, space_after_pt=4)
    r = title.add_run("Бережливая канва проекта (Lean Canvas)")
    apply_run_font(r, 14, bold=True)

    subtitle = doc.add_paragraph()
    set_paragraph_format(subtitle, align=WD_ALIGN_PARAGRAPH.CENTER, first_indent_cm=None, line_spacing=1.0, space_before_pt=0, space_after_pt=6)
    r = subtitle.add_run(PROJECT_TITLE)
    apply_run_font(r, 11)

    meta = doc.add_paragraph()
    set_paragraph_format(meta, align=WD_ALIGN_PARAGRAPH.CENTER, first_indent_cm=None, line_spacing=1.0, space_before_pt=0, space_after_pt=8)
    r = meta.add_run(f"Автор: {PARTICIPANT} | {UNIVERSITY} | Статус: функциональный прототип, подготовка к пилотной апробации")
    apply_run_font(r, 9)

    table = doc.add_table(rows=3, cols=6)
    table.style = "Table Grid"

    for row in table.rows:
        for cell in row.cells:
            cell.width = Cm(4.35)

    problem = table.cell(0, 0).merge(table.cell(1, 0))
    solution = table.cell(0, 1)
    uvp = table.cell(0, 2).merge(table.cell(1, 3))
    unfair = table.cell(0, 4)
    segments = table.cell(0, 5).merge(table.cell(1, 5))
    metrics = table.cell(1, 1)
    channels = table.cell(1, 4)
    costs = table.cell(2, 0).merge(table.cell(2, 2))
    revenue = table.cell(2, 3).merge(table.cell(2, 5))

    add_canvas_block(
        problem,
        "2. Проблемы",
        [
            "Учёт объектов, обсуждения, фото и документы разнесены между несколькими несвязанными инструментами.",
            "При нестабильной связи централизованный контур становится недоступным, и работа уходит в ручной режим.",
            "После выездной инвентаризации результаты приходится повторно сводить и переносить вручную.",
            "Существующие альтернативы: централизованные учётные системы, Excel, бумажные формы, внешние мессенджеры, отдельные решения для маркировки и сканирования.",
        ],
        fill="F8FAFC",
    )
    add_canvas_block(
        solution,
        "4. Решение",
        [
            "Гибридная архитектура: локальные клиентские узлы плюс центральный контур консолидации.",
            "Карточки имущества, сессии, замечания, инциденты, маркировка и сканирование в одном рабочем контуре.",
            "Встроенные чаты, обсуждения, файловый обмен, звонки и последующая синхронизация после восстановления связи.",
        ],
        fill="F8FAFC",
    )
    add_canvas_block(
        uvp,
        "3. Уникальное ценностное предложение",
        [
            "Единый рабочий контур для инвентаризации и служебного взаимодействия, сохраняющий работоспособность даже при потере устойчивой связи.",
            "Инвентаризация, коммуникации и последующая консолидация данных не разорваны между разными системами.",
            "Решение ориентировано не на офисный сценарий, а на удалённые объекты, где связь ограничена или нестабильна.",
        ],
        subtitle="Высокоуровневая концепция: инвентаризация + взаимодействие сотрудников + синхронизация в одной системе",
        fill="EAF3FF",
    )
    add_canvas_block(
        unfair,
        "9. Несправедливое преимущество",
        [
            "Сформированный инженерный задел: настольный и Android-клиенты, документация, схемы и тесты.",
            "Единый контракт MeshNode и подтверждённые сценарии автономной работы, межузлового обмена и последующей синхронизации.",
            "Готовность к университетской апробации как к контролируемому пилотному сценарию.",
        ],
        fill="F8FAFC",
    )
    add_canvas_block(
        segments,
        "1. Сегменты пользователей",
        [
            "Инвентаризационные комиссии и сотрудники удалённых объектов.",
            "Территориально распределённые организации, склады, базы снабжения, месторождения, лабораторные и учебные комплексы.",
            "Ранние пользователи: вузовская пилотная площадка, подразделения с локальной сетью и периодически недоступным внешним каналом.",
        ],
        fill="F8FAFC",
    )
    add_canvas_block(
        metrics,
        "8. Ключевые метрики",
        [
            "Доля операций, выполненных без доступа к центральному контуру.",
            "Успешность синхронизации накопленных изменений и количество конфликтов.",
            "Время проведения инвентаризационной сессии и число ручных переносов после выезда.",
            "Полнота заполнения карточек, замечаний, вложений и инцидентов непосредственно в рабочем контуре.",
        ],
        fill="F8FAFC",
    )
    add_canvas_block(
        channels,
        "5. Каналы выхода",
        [
            "Студенческая олимпиада «Газпром» и профильные инженерные площадки.",
            "Университетская апробация в НИЯУ МИФИ.",
            "Пилотные демонстрации на ограниченных объектах и прямой диалог с организациями, для которых критична устойчивая работа при ограниченной связности.",
        ],
        fill="F8FAFC",
    )
    add_canvas_block(
        costs,
        "7. Структура затрат",
        [
            "Разработка и поддержка клиентского и центрального контуров.",
            "Тестирование сетевых сценариев, синхронизации и устойчивости обмена.",
            "Инфраструктура центрального контура, хранение вложений и артефактов выгрузки.",
            "Пилотная апробация, адаптация под организационные регламенты и сопровождение внедрения.",
        ],
        fill="F8FAFC",
    )
    add_canvas_block(
        revenue,
        "6. Источники ценности и модели монетизации",
        [
            "На текущем этапе проект находится на стадии функционального прототипа и пилотной проработки.",
            "В качестве рабочих вариантов рассматриваются лицензирование программного решения, проектная адаптация и внедрение.",
            "Дополнительные варианты: техническое сопровождение, доработки под инфраструктуру заказчика и интеграция с его централизованными сервисами.",
        ],
        fill="F8FAFC",
    )

    note = doc.add_paragraph()
    set_paragraph_format(note, align=WD_ALIGN_PARAGRAPH.LEFT, first_indent_cm=1.25, line_spacing=1.15, space_before_pt=8, space_after_pt=0)
    r = note.add_run(
        "Примечание. Блоки, связанные с каналами вывода и монетизацией, отражают рабочие гипотезы пилотного и внедренческого развития проекта и не трактуются как уже реализованная коммерческая модель."
    )
    apply_run_font(r, 9)

    doc.save(FINAL_FILES["lean_canvas"])


def build_cover_letter_doc() -> None:
    doc = Document()
    setup_doc_styles(doc)
    setup_footer(doc)
    p = doc.add_paragraph("Сопроводительное письмо", style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph()

    rec = doc.add_paragraph()
    rec.alignment = WD_ALIGN_PARAGRAPH.LEFT
    rec.add_run("В оргкомитет Студенческой олимпиады «Газпром»").bold = True

    doc.add_paragraph()
    body = [
        "Направляю заявку на участие в конкурсе проектов заключительного этапа по профилю «Информационные системы и технологии», а также пояснительную записку по проекту:",
        f"«{PROJECT_TITLE}».",
        "Проект посвящён разработке прикладной распределённой информационной системы для учёта имущества, работы инвентаризационных комиссий и служебного взаимодействия сотрудников на удалённых и территориально распределённых объектах в условиях нестабильной связи.",
        "Во вложении направляю заявку на конкурс проектов, пояснительную записку и презентационные материалы в установленные сроки.",
    ]
    for paragraph in body:
        doc.add_paragraph(paragraph)

    doc.add_paragraph("Контактные данные:")
    for item in [PARTICIPANT, UNIVERSITY, "tigor7750@gmail.com"]:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_paragraph()
    sig = doc.add_paragraph("С уважением,\nТитарь Игорь Андреевич")
    sig.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    doc.save(FINAL_FILES["cover"])


def build_appendices_doc(assets: dict[str, Path]) -> None:
    doc = Document()
    setup_doc_styles(doc)
    setup_footer(doc)
    add_title_page(doc, "Приложения, схемы и иллюстрации", [PROJECT_TITLE, PARTICIPANT])

    intro = doc.add_paragraph("В приложениях приведены основные схемы, используемые в заявке, пояснительной записке и презентационных материалах.")
    intro.paragraph_format.space_after = Pt(12)

    figures = [
        ("Приложение А. Компонентная архитектура", assets["architecture_existing"], "Рисунок А.1. Компонентная архитектура проекта. Источник: техническая презентация проекта.", 17.0),
        ("Приложение Б. Режимы работы при наличии и отсутствии внешней связи", assets["online_offline"], "Рисунок Б.1. Режимы работы при наличии и отсутствии внешней связи. Источник: подготовлено автором на основе архитектурных материалов проекта.", 17.0),
        ("Приложение В. Синхронизация изменений", assets["sync"], "Рисунок В.1. Локальная очередь изменений и последующая синхронизация с центральным контуром. Источник: подготовлено автором на основе архитектурных материалов проекта.", 17.0),
        ("Приложение Г. Сценарий работы комиссии", assets["commission"], "Рисунок Г.1. Типовой сценарий работы комиссии на удалённом объекте. Источник: подготовлено автором на основе проектных материалов.", 17.0),
        ("Приложение Д. Верхнеуровневая ER-модель инвентаризации", assets["inventory_er_existing"], "Рисунок Д.1. Верхнеуровневая ER-модель инвентаризационного контура. Источник: техническая презентация проекта.", 17.0),
        ("Приложение Е. Верхнеуровневая ER-модель коммуникаций и синхронизации", assets["comm_er_existing"], "Рисунок Е.1. Верхнеуровневая ER-модель мессенджера, файлов, звонков и синхронизации. Источник: техническая презентация проекта.", 17.0),
        ("Приложение Ж. Центральный контур: сервисы и взаимодействие", assets["central_contour_services"], "Рисунок Ж.1. Схема сервисов и взаимодействий в центральном контуре. Источник: подготовлено автором на основе архитектурных материалов проекта и expert-link-core.", 17.0),
    ]
    for heading, img_path, caption, width in figures:
        doc.add_paragraph(heading, style="Heading 1")
        insert_figure(doc, img_path, caption, width)
        doc.add_paragraph()

    doc.add_paragraph("Приложение З. Реальные экраны прототипа", style="Heading 1")
    for title, key, caption in [
        ("Настольный клиент: главный экран", "proto_main", "Рисунок З.1. Главный экран настольного прототипа. Источник: техническая презентация проекта."),
        ("Настольный клиент: сводная панель оборудования", "proto_inventory_table", "Рисунок З.2. Табличный режим сводной панели оборудования. Источник: техническая презентация проекта."),
        ("Настольный клиент: фильтры и поиск", "proto_inventory_filters", "Рисунок З.3. Расширенный поиск и фильтрация. Источник: техническая презентация проекта."),
        ("Настольный клиент: правка объекта", "proto_inventory_edit", "Рисунок З.4. Окно правки карточки объекта. Источник: техническая презентация проекта."),
        ("Мобильный клиент: главный экран", "mobile_home", "Рисунок З.5. Главный экран мобильного клиента. Источник: материалы проекта."),
        ("Мобильный клиент: чаты", "mobile_chats", "Рисунок З.6. Экран чатов мобильного клиента. Источник: материалы проекта."),
        ("Мобильный клиент: передача файлов", "mobile_transfers", "Рисунок З.7. Экран передачи файлов мобильного клиента. Источник: материалы проекта."),
        ("Мобильный клиент: звонок", "mobile_call", "Рисунок З.8. Экран звонка мобильного клиента. Источник: материалы проекта."),
    ]:
        doc.add_paragraph(title, style="Heading 2")
        insert_figure(doc, assets[key], caption, 14.5 if key.startswith("mobile") else 17.0)
        doc.add_paragraph()

    doc.save(FINAL_FILES["appendices"])


def build_sources_doc() -> None:
    build_doc_from_md(
        "08-sources.md",
        FINAL_FILES["sources"],
        "Использованные источники",
        [PROJECT_TITLE, UNIVERSITY, PARTICIPANT],
        title_page=True,
        manual_contents=False,
    )


def build_checklist_doc() -> None:
    doc = Document()
    setup_doc_styles(doc)
    setup_footer(doc)
    p = doc.add_paragraph("Что проверить перед отправкой", style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph()
    doc.add_paragraph("Ручная проверка данных:")
    for item in [
        "Площадка защиты в основной заявке.",
        "Официальная формулировка вуза участника.",
        "ФИО, электронная почта и при необходимости телефон.",
        "Совпадение реквизитов участника с данными в личном кабинете олимпиады.",
    ]:
        doc.add_paragraph(item, style="List Bullet")
    doc.add_paragraph()
    doc.add_paragraph("Файлы финального комплекта:")
    for path in FINAL_FILES.values():
        if path.suffix in {".docx", ".pptx"}:
            doc.add_paragraph(path.name, style="List Bullet")
    doc.add_paragraph()
    doc.add_paragraph("Техническая проверка перед отправкой:")
    for item in [
        "Открываются ли все DOCX и PPTX в LibreOffice или Microsoft Office без ошибок.",
        "Не остались ли в финальных файлах служебные поля, домашние пути и markdown-разметка.",
        "Нет ли в презентациях акселераторских слайдов про инвестиционный запрос.",
        "Совпадает ли название проекта во всех документах и на титульных слайдах.",
    ]:
        doc.add_paragraph(item, style="List Bullet")
    doc.save(FINAL_FILES["checklist"])


def copy_delivery_aliases() -> None:
    shutil.copy2(FINAL_FILES["application"], DELIVERY_ALIASES["application"])
    shutil.copy2(FINAL_FILES["note"], DELIVERY_ALIASES["note"])


def create_prs() -> Presentation:
    prs = Presentation()
    prs.slide_width = PptInches(13.333)
    prs.slide_height = PptInches(7.5)
    return prs


def add_slide_bg(slide, kind: str = "business") -> None:
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = PPT_WHITE
    bar = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, 0, prs_w(slide, 13.333), PptInches(0.16))
    bar.fill.solid()
    bar.fill.fore_color.rgb = PPT_BLUE if kind == "technical" else PPT_TEAL
    bar.line.fill.background()
    for x, y, w, h in [(-0.5, -0.5, 2.8, 2.8), (11.1, -0.3, 2.7, 2.7), (10.8, 6.0, 3.0, 3.0)]:
        shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, PptInches(x), PptInches(y), PptInches(w), PptInches(h))
        shape.fill.background()
        shape.line.color.rgb = PPT_LIGHT_GRAY
        shape.line.width = PptPt(1.25)


def prs_w(slide, inches: float):
    return PptInches(inches)


def add_footer(slide, idx: int, label: str) -> None:
    left = slide.shapes.add_textbox(PptInches(0.5), PptInches(7.1), PptInches(6.0), PptInches(0.25))
    tf = left.text_frame
    tf.text = label
    p = tf.paragraphs[0]
    p.font.name = "Arial"
    p.font.size = PptPt(10)
    p.font.color.rgb = PPT_MID_GRAY

    right = slide.shapes.add_textbox(PptInches(12.2), PptInches(7.05), PptInches(0.5), PptInches(0.3))
    tf = right.text_frame
    tf.text = str(idx)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    p.font.name = "Arial"
    p.font.size = PptPt(10)
    p.font.color.rgb = PPT_MID_GRAY


def add_title_text(slide, title: str, subtitle: str | None = None) -> None:
    tb = slide.shapes.add_textbox(PptInches(0.6), PptInches(0.55), PptInches(8.2), PptInches(1.1))
    tf = tb.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = title
    p.font.name = "Arial"
    p.font.size = PptPt(24)
    p.font.bold = True
    p.font.color.rgb = PPT_BLUE
    if subtitle:
        p2 = tf.add_paragraph()
        p2.text = subtitle
        p2.font.name = "Arial"
        p2.font.size = PptPt(11)
        p2.font.color.rgb = PPT_MID_GRAY
        p2.space_before = PptPt(4)


def add_body_bullets(slide, bullets: list[str], x: float, y: float, w: float, h: float, font_size: int = 18) -> None:
    tb = slide.shapes.add_textbox(PptInches(x), PptInches(y), PptInches(w), PptInches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.clear()
    for idx, bullet in enumerate(bullets):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = bullet
        p.level = 0
        p.bullet = True
        p.font.name = "Arial"
        p.font.size = PptPt(font_size)
        p.font.color.rgb = PPT_DARK
        p.space_after = PptPt(8)


def add_text_block(slide, text: str, x: float, y: float, w: float, h: float, font_size: int = 18, color: RGBColor = PPT_DARK, bold: bool = False, align: PP_ALIGN = PP_ALIGN.LEFT) -> None:
    tb = slide.shapes.add_textbox(PptInches(x), PptInches(y), PptInches(w), PptInches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = align
    p.font.name = "Arial"
    p.font.size = PptPt(font_size)
    p.font.bold = bold
    p.font.color.rgb = color


def add_picture(slide, path: Path, x: float, y: float, w: float | None = None, h: float | None = None) -> None:
    kwargs = {}
    if w is not None:
        kwargs["width"] = PptInches(w)
    if h is not None:
        kwargs["height"] = PptInches(h)
    slide.shapes.add_picture(str(path), PptInches(x), PptInches(y), **kwargs)


def add_card(slide, x: float, y: float, w: float, h: float, title: str, subtitle: str | None = None, fill: RGBColor = PPT_LIGHT_BLUE, title_color: RGBColor = PPT_BLUE) -> None:
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, PptInches(x), PptInches(y), PptInches(w), PptInches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = PPT_BORDER
    shape.line.width = PptPt(1)
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title
    p.font.name = "Arial"
    p.font.size = PptPt(18)
    p.font.bold = True
    p.font.color.rgb = title_color
    if subtitle:
        p2 = tf.add_paragraph()
        p2.text = subtitle
        p2.font.name = "Arial"
        p2.font.size = PptPt(12)
        p2.font.color.rgb = PPT_DARK
        p2.space_before = PptPt(6)


def add_title_slide(prs: Presentation, heading: str, subheading: str, kind: str) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_slide_bg(slide, kind)
    banner = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, PptInches(0.6), PptInches(1.2), PptInches(12.0), PptInches(4.35))
    banner.fill.solid()
    banner.fill.fore_color.rgb = PPT_LIGHT_BLUE if kind == "technical" else PPT_WHITE
    banner.line.color.rgb = PPT_BORDER
    banner.line.width = PptPt(1.2)
    add_text_block(slide, heading, 0.95, 1.62, 8.25, 3.1, font_size=24, color=PPT_BLUE, bold=True)
    add_text_block(slide, subheading, 0.98, 4.75, 7.25, 0.6, font_size=14, color=PPT_DARK)
    add_card(slide, 9.45, 1.65, 2.5, 1.05, "Участник", PARTICIPANT, fill=PPT_LIGHT_BLUE)
    add_card(slide, 9.45, 2.95, 2.5, 1.05, "Профиль", "Информационные системы и технологии", fill=PPT_LIGHT_BLUE)
    add_card(slide, 9.45, 4.25, 2.5, 1.05, "Вуз", UNIVERSITY, fill=PPT_LIGHT_BLUE)
    add_footer(slide, len(prs.slides), PROJECT_SHORT)


def build_business_pptx(assets: dict[str, Path]) -> None:
    prs = create_prs()
    add_title_slide(prs, PROJECT_TITLE, "Бизнес-презентация для конкурсной подачи на олимпиаду «Газпром»", "business")

    agenda = prs.slides.add_slide(prs.slide_layouts[6])
    add_slide_bg(agenda, "business")
    add_title_text(agenda, "Структура презентации")
    for idx, item in enumerate([
        "Проблема удалённых объектов",
        "Предлагаемое решение",
        "Сценарии применения",
        "Сравнение с классическими системами",
        "Готовность к пилоту",
        "Практическая ценность для Газпром",
    ]):
        add_card(agenda, 0.9 + (idx % 2) * 5.9, 1.8 + (idx // 2) * 1.45, 5.3, 1.0, f"{idx + 1}. {item}", fill=PPT_WHITE if idx % 2 else PPT_LIGHT_BLUE)
    add_footer(agenda, len(prs.slides), PROJECT_SHORT)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_slide_bg(slide, "business")
    add_title_text(slide, "Почему задача важна для удалённых объектов")
    add_body_bullets(slide, [
        "Учёт имущества, материалы проверки и обсуждения часто живут в разных каналах.",
        "На удалённых площадках внешний канал может быть нестабилен или отсутствовать полностью.",
        "Комиссия вынуждена переходить на таблицы, бумагу и сторонние мессенджеры.",
        "После этого результат приходится вручную сводить и переносить в централизованный контур.",
    ], 0.8, 1.8, 5.8, 4.5, 18)
    add_picture(slide, assets["fragmentation"], 6.8, 1.55, 5.6)
    add_footer(slide, len(prs.slides), PROJECT_SHORT)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_slide_bg(slide, "business")
    add_title_text(slide, "Что именно нарушается в условиях нестабильной связи")
    add_body_bullets(slide, [
        "невозможно опираться только на центральную систему;",
        "теряется контекст спорных объектов и принятых решений;",
        "обсуждения и файлы не связаны с карточкой имущества;",
        "проверка на площадке замедляется, а результаты откладываются «до офиса».",
    ], 0.9, 1.85, 6.0, 4.2, 19)
    add_card(slide, 0.9, 5.55, 6.0, 0.95, "Следствие", "Потеря времени, рост числа ручных переносов и снижение прозрачности учёта.", fill=PPT_LIGHT_BLUE)
    add_picture(slide, assets["mobile_home"], 8.4, 1.8, 2.1)
    add_picture(slide, assets["mobile_profile"], 10.7, 1.8, 2.1)
    add_footer(slide, len(prs.slides), PROJECT_SHORT)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_slide_bg(slide, "business")
    add_title_text(slide, "Что предлагает проект")
    add_picture(slide, assets["single_contour"], 0.8, 1.6, 7.2)
    add_body_bullets(slide, [
        "единый контур для учёта имущества и оперативного взаимодействия;",
        "локальная работа между устройствами без обязательной зависимости от центрального интернета;",
        "последующая синхронизация с центральным контуром при восстановлении связи;",
    ], 8.35, 2.0, 4.1, 3.2, 18)
    add_card(slide, 8.35, 5.45, 4.05, 0.9, "Ключевая идея", "Не отдельный мессенджер и не отдельный учётный модуль, а единая распределённая платформа.", fill=PPT_LIGHT_BLUE)
    add_footer(slide, len(prs.slides), PROJECT_SHORT)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_slide_bg(slide, "business")
    add_title_text(slide, "Как система работает с интернетом и без него")
    add_picture(slide, assets["online_offline"], 0.8, 1.6, 7.25)
    add_body_bullets(slide, [
        "в online-режиме данные консолидируются в центральном контуре;",
        "в offline-режиме комиссия продолжает работу локально;",
        "пользователь не теряет доступ к карточкам, сессиям, чатам, файлам и обсуждениям;",
    ], 8.3, 2.05, 4.1, 3.2, 17)
    add_footer(slide, len(prs.slides), PROJECT_SHORT)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_slide_bg(slide, "business")
    add_title_text(slide, "Ключевые функции платформы")
    add_card(slide, 0.9, 1.8, 2.8, 1.2, "Инвентаризация", "Карточки объектов, статусы, местоположения, ответственные", fill=PPT_LIGHT_BLUE)
    add_card(slide, 3.95, 1.8, 2.8, 1.2, "Маркировка", "QR, штрихкоды, печать этикеток, сканирование", fill=PPT_LIGHT_BLUE)
    add_card(slide, 0.9, 3.25, 2.8, 1.2, "Коммуникации", "Чаты, ветки обсуждений, файлы, звонки", fill=PPT_LIGHT_BLUE)
    add_card(slide, 3.95, 3.25, 2.8, 1.2, "Синхронизация", "Локальная очередь изменений и central sync", fill=PPT_LIGHT_BLUE)
    add_picture(slide, assets["proto_inventory_table"], 7.15, 1.7, 5.2)
    add_footer(slide, len(prs.slides), PROJECT_SHORT)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_slide_bg(slide, "business")
    add_title_text(slide, "Сценарий работы комиссии")
    add_picture(slide, assets["commission"], 0.8, 1.75, 7.2)
    add_body_bullets(slide, [
        "подготовка сессии и списка объектов;",
        "выезд на площадку и сканирование оборудования;",
        "фиксация замечаний, фото и документов;",
        "обсуждение спорных объектов внутри рабочего контура;",
        "синхронизация после восстановления связи;",
    ], 8.25, 1.95, 4.1, 3.8, 17)
    add_footer(slide, len(prs.slides), PROJECT_SHORT)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_slide_bg(slide, "business")
    add_title_text(slide, "Почему это лучше разрозненного набора инструментов")
    add_card(slide, 0.9, 1.9, 2.7, 1.15, "Excel + бумага", "Данные и промежуточные результаты живут отдельно от обсуждений", fill=PPT_WHITE)
    add_card(slide, 0.9, 3.25, 2.7, 1.15, "Сторонний мессенджер", "Файлы и решения комиссии не привязаны к объектам", fill=PPT_WHITE)
    add_card(slide, 0.9, 4.6, 2.7, 1.15, "Только central-контур", "При потере связи часть работы откладывается", fill=PPT_WHITE)
    add_card(slide, 4.05, 2.55, 4.15, 2.1, "Платформа проекта", "Объекты, проверки, коммуникации и материалы объединены в одном рабочем контуре. Отсутствие внешней связи не останавливает локальную работу.", fill=PPT_LIGHT_BLUE)
    add_picture(slide, assets["mobile_chats"], 9.1, 1.85, 1.9)
    add_picture(slide, assets["mobile_transfers"], 11.15, 1.85, 1.9)
    add_footer(slide, len(prs.slides), PROJECT_SHORT)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_slide_bg(slide, "business")
    add_title_text(slide, "Сравнение с классическими централизованными решениями")
    headers = ["Критерий", "Наш проект", "Классические централизованные решения"]
    rows = [
        ("Локальная автономная работа без центрального канала", "Да", "Как правило, не является базовым сценарием"),
        ("Встроенные коммуникации в предметном контуре", "Да", "Обычно используются внешние каналы"),
        ("Связка объекта, обсуждения и вложений", "Да", "Часто требует ручного связывания"),
        ("Работа комиссии на удалённой площадке", "Целевой сценарий", "Обычно опирается на устойчивый доступ к центру"),
    ]
    table = slide.shapes.add_table(len(rows) + 1, len(headers), PptInches(0.8), PptInches(1.85), PptInches(11.8), PptInches(3.8)).table
    for idx, header in enumerate(headers):
        cell = table.cell(0, idx)
        cell.text = header
        cell.fill.solid()
        cell.fill.fore_color.rgb = PPT_LIGHT_BLUE
    for r_idx, row in enumerate(rows, start=1):
        for c_idx, value in enumerate(row):
            table.cell(r_idx, c_idx).text = value
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.text_frame.paragraphs:
                paragraph.font.name = "Arial"
                paragraph.font.size = PptPt(14)
                paragraph.font.color.rgb = PPT_DARK
    add_footer(slide, len(prs.slides), PROJECT_SHORT)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_slide_bg(slide, "business")
    add_title_text(slide, "Текущий статус и готовность к пилоту")
    add_card(slide, 0.9, 1.95, 3.1, 1.15, "Статус", "Функциональный прототип", fill=PPT_LIGHT_BLUE)
    add_card(slide, 0.9, 3.35, 3.1, 1.15, "Подтверждено", "API, клиентские экраны, inventory-контур, hybrid sync, тесты", fill=PPT_LIGHT_BLUE)
    add_card(slide, 0.9, 4.75, 3.1, 1.15, "Следующий шаг", "Пилотная апробация на ограниченной площадке", fill=PPT_LIGHT_BLUE)
    add_picture(slide, assets["proto_main"], 4.45, 1.8, 4.0)
    add_picture(slide, assets["mobile_pairing_scan"], 9.0, 1.85, 1.5)
    add_picture(slide, assets["mobile_call"], 10.75, 1.85, 1.5)
    add_footer(slide, len(prs.slides), PROJECT_SHORT)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_slide_bg(slide, "business")
    add_title_text(slide, "План пилотного внедрения")
    phases = [
        ("Этап 1", "Одна площадка", "реестр объектов, маркировка, одна-две сессии"),
        ("Этап 2", "Проверка offline-сценариев", "локальная работа нескольких узлов и последующая синхронизация"),
        ("Этап 3", "Расширение и интеграция", "увеличение числа пользователей и подготовка к подключению центрального контура"),
    ]
    for idx, (phase, title, desc) in enumerate(phases):
        add_card(slide, 0.9 + idx * 4.1, 2.1, 3.5, 2.25, f"{phase}: {title}", desc, fill=PPT_LIGHT_BLUE if idx == 1 else PPT_WHITE)
    add_card(slide, 2.0, 5.1, 9.6, 0.95, "Результат", "Пилот позволяет подтвердить удобство сценариев, сетевую устойчивость и применимость решения к распределённым объектам.", fill=PPT_LIGHT_BLUE)
    add_footer(slide, len(prs.slides), PROJECT_SHORT)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_slide_bg(slide, "business")
    add_title_text(slide, "Практическая ценность для удалённых производственных объектов")
    add_card(slide, 0.9, 2.0, 2.8, 1.2, "Непрерывность работы", "Проверка и взаимодействие не останавливаются при потере внешней связи", fill=PPT_LIGHT_BLUE)
    add_card(slide, 3.95, 2.0, 2.8, 1.2, "Единый контур", "Карточка объекта, обсуждение, файл и результат связаны между собой", fill=PPT_LIGHT_BLUE)
    add_card(slide, 7.0, 2.0, 2.8, 1.2, "Прозрачность учёта", "Меньше ручных переносов и разрывов между фактом и отчётностью", fill=PPT_LIGHT_BLUE)
    add_card(slide, 2.45, 3.65, 2.8, 1.2, "Готовность к пилоту", "Есть прототип и набор сценариев для ограниченной апробации", fill=PPT_LIGHT_BLUE)
    add_card(slide, 5.55, 3.65, 2.8, 1.2, "Релевантность профилю ИСТ", "Прикладная распределённая информационная система для реального процесса", fill=PPT_LIGHT_BLUE)
    add_footer(slide, len(prs.slides), PROJECT_SHORT)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_slide_bg(slide, "business")
    add_title_text(slide, "Вывод")
    add_card(slide, 0.9, 1.9, 11.5, 3.05, "Распределённая платформа инвентаризации позволяет не останавливать работу комиссии и сотрудников там, где централизованный интернет нестабилен или отсутствует. Проект объединяет учёт, маркировку, обсуждения, файлы и связь в одном рабочем контуре и уже находится на стадии функционального прототипа, пригодного для пилотной апробации.", fill=PPT_LIGHT_BLUE, title_color=PPT_DARK)
    add_text_block(slide, f"{PARTICIPANT}\n{UNIVERSITY}\ntigor7750@gmail.com", 0.95, 5.45, 4.6, 0.9, font_size=16, color=PPT_DARK)
    add_footer(slide, len(prs.slides), PROJECT_SHORT)

    prs.save(FINAL_FILES["business_pptx"])


def build_technical_pptx(assets: dict[str, Path]) -> None:
    prs = create_prs()
    add_title_slide(prs, PROJECT_TITLE, "Техническая презентация для конкурсной защиты", "technical")

    agenda = prs.slides.add_slide(prs.slide_layouts[6])
    add_slide_bg(agenda, "technical")
    add_title_text(agenda, "Техническая структура")
    for idx, item in enumerate([
        "Ограничения среды и сценарии использования",
        "Архитектура и контуры системы",
        "Online/offline и синхронизация",
        "Сетевое взаимодействие и коммуникации",
        "Инвентаризационный контур и прототип",
        "Статус реализации и ограничения",
    ]):
        add_card(agenda, 0.9 + (idx % 2) * 5.9, 1.8 + (idx // 2) * 1.45, 5.35, 1.0, f"{idx + 1}. {item}", fill=PPT_LIGHT_BLUE if idx % 2 == 0 else PPT_WHITE)
    add_footer(agenda, len(prs.slides), PROJECT_SHORT)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_slide_bg(slide, "technical")
    add_title_text(slide, "Ограничения среды и сценарии использования")
    add_body_bullets(slide, [
        "удалённые объекты и территориально распределённые площадки;",
        "нестабильная или отсутствующая связь с центральным контуром;",
        "необходимость продолжать учёт и взаимодействие без внешнего канала;",
        "последующая консолидация данных после восстановления связи;",
    ], 0.85, 1.85, 6.0, 4.2, 18)
    add_picture(slide, assets["fragmentation"], 7.1, 1.6, 5.1)
    add_footer(slide, len(prs.slides), PROJECT_SHORT)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_slide_bg(slide, "technical")
    add_title_text(slide, "Обзор решения")
    add_picture(slide, assets["single_contour"], 0.75, 1.7, 6.9)
    add_body_bullets(slide, [
        "edge-узел сочетает inventory-контур, локальное хранилище и mesh-коммуникации;",
        "central-контур используется как канонический источник и точка консолидации;",
        "локальные операции не блокируются недоступностью центральной системы;",
    ], 7.8, 2.0, 4.1, 3.0, 17)
    add_footer(slide, len(prs.slides), PROJECT_SHORT)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_slide_bg(slide, "technical")
    add_title_text(slide, "Компонентная архитектура")
    add_picture(slide, assets["architecture_existing"], 0.75, 1.55, 8.0)
    add_body_bullets(slide, [
        "клиентские устройства: desktop и Android;",
        "локальный контур: UI, inventory, локальная БД, очередь синхронизации, mesh-модуль;",
        "central-контур: доступ, inventory, sync, вложения и экспорт;",
    ], 9.05, 1.95, 3.3, 3.1, 16)
    add_footer(slide, len(prs.slides), PROJECT_SHORT)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_slide_bg(slide, "technical")
    add_title_text(slide, "Edge-контур узла")
    add_card(slide, 0.85, 1.9, 3.0, 1.1, "UI и workflow", "Карточки имущества, сессии, scanner, отчёты", fill=PPT_LIGHT_BLUE)
    add_card(slide, 0.85, 3.25, 3.0, 1.1, "Локальное хранение", "Состояние узла, данные предметного контура, очередь изменений", fill=PPT_LIGHT_BLUE)
    add_card(slide, 0.85, 4.6, 3.0, 1.1, "Коммуникации", "Discovery, чаты, файлы, звонки, topology и diagnostics", fill=PPT_LIGHT_BLUE)
    add_picture(slide, assets["proto_main"], 4.25, 1.7, 7.8)
    add_footer(slide, len(prs.slides), PROJECT_SHORT)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_slide_bg(slide, "technical")
    add_title_text(slide, "Online/offline режимы")
    add_picture(slide, assets["online_offline"], 0.8, 1.65, 7.4)
    add_body_bullets(slide, [
        "online: local-first работа + central sync;",
        "offline: локальная автономность и взаимодействие между соседними узлами;",
        "в обоих режимах доступны предметные сценарии, обсуждения и файлы;",
    ], 8.35, 2.0, 4.0, 3.0, 17)
    add_footer(slide, len(prs.slides), PROJECT_SHORT)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_slide_bg(slide, "technical")
    add_title_text(slide, "Верхнеуровневая ER-модель инвентаризации")
    add_picture(slide, assets["inventory_er_existing"], 0.85, 1.55, 8.2)
    add_body_bullets(slide, [
        "ключевые сущности: организация, объект, сессия, review, incident, attachment, code;",
        "отдельно поддержаны маркировка, печать и история событий;",
        "модель рассчитана на инвентаризационные сессии и рабочие роли участников;",
    ], 9.2, 2.0, 3.1, 3.2, 16)
    add_footer(slide, len(prs.slides), PROJECT_SHORT)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_slide_bg(slide, "technical")
    add_title_text(slide, "ER-модель коммуникаций и синхронизации")
    add_picture(slide, assets["comm_er_existing"], 0.85, 1.55, 8.2)
    add_body_bullets(slide, [
        "conversation, message и thread образуют discussion-контур;",
        "file transfer и call session интегрированы в тот же рабочий контур;",
        "peer device, sync state и sync change фиксируют состояние обмена между узлами;",
    ], 9.2, 2.0, 3.1, 3.2, 16)
    add_footer(slide, len(prs.slides), PROJECT_SHORT)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_slide_bg(slide, "technical")
    add_title_text(slide, "Discovery, pairing и доверие")
    add_body_bullets(slide, [
        "локальное обнаружение соседних узлов;",
        "invite с ограничением по времени и nonce;",
        "peerId как криптографическая идентичность узла;",
        "trusted state после успешного pairing;",
    ], 0.9, 1.95, 5.9, 3.6, 18)
    add_picture(slide, assets["mobile_pairing_scan"], 7.25, 1.8, 2.2)
    add_picture(slide, assets["desktop_profile"], 9.65, 1.8, 2.45)
    add_footer(slide, len(prs.slides), PROJECT_SHORT)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_slide_bg(slide, "technical")
    add_title_text(slide, "Обработка входящего сетевого пакета")
    add_picture(slide, assets["packet_sequence_existing"], 0.85, 1.55, 7.4)
    add_body_bullets(slide, [
        "controller принимает транспортную оболочку пакета;",
        "lifecycle выполняет dedup, trust и signature validation;",
        "payload передаётся в соответствующий use-case сервис;",
        "по результату формируется ответ доставки;",
    ], 8.5, 1.95, 3.8, 3.5, 16)
    add_footer(slide, len(prs.slides), PROJECT_SHORT)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_slide_bg(slide, "technical")
    add_title_text(slide, "Передача файлов")
    add_picture(slide, assets["file_flow_existing"], 0.85, 1.55, 7.4)
    add_body_bullets(slide, [
        "offer/accept для начала сессии передачи;",
        "чанки фиксированного размера с подтверждением FILE_ACK;",
        "завершение по FILE_COMPLETE и дозагрузка по FILE_RESUME_REQUEST;",
    ], 8.45, 2.0, 3.8, 3.0, 16)
    add_picture(slide, assets["mobile_transfers"], 9.15, 4.8, 1.85)
    add_footer(slide, len(prs.slides), PROJECT_SHORT)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_slide_bg(slide, "technical")
    add_title_text(slide, "Звонки и media-контур")
    add_picture(slide, assets["call_flow_existing"], 0.8, 1.5, 7.4)
    add_body_bullets(slide, [
        "signaling идёт внутри рабочего контура через CALL_INVITE и CALL_SIGNAL;",
        "SDP/ICE используются для согласования WebRTC-сессии;",
        "в материалах проекта подтверждены Android- и desktop-клиенты с media backend;",
    ], 8.4, 1.95, 3.9, 3.0, 16)
    add_picture(slide, assets["mobile_call"], 9.25, 4.85, 1.8)
    add_footer(slide, len(prs.slides), PROJECT_SHORT)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_slide_bg(slide, "technical")
    add_title_text(slide, "Инвентаризационный сценарий на удалённом объекте")
    add_picture(slide, assets["commission"], 0.85, 1.75, 7.2)
    add_body_bullets(slide, [
        "сканирование объекта открывает карточку и действия по проверке;",
        "замечания, фото и комментарии можно привязать к объекту сразу на месте;",
        "обсуждение спорных случаев ведётся в связанном коммуникационном контуре;",
    ], 8.25, 2.0, 4.1, 3.2, 16)
    add_picture(slide, assets["proto_inventory_filters"], 9.25, 4.85, 2.7)
    add_footer(slide, len(prs.slides), PROJECT_SHORT)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_slide_bg(slide, "technical")
    add_title_text(slide, "Синхронизация изменений")
    add_picture(slide, assets["sync"], 0.85, 1.65, 7.45)
    add_body_bullets(slide, [
        "локальные изменения сначала фиксируются у пользователя и узла;",
        "pending queue отделяет локальную работу от central upload/pull;",
        "конфликты не теряются и остаются видимыми для последующего разбора;",
    ], 8.45, 2.0, 3.9, 3.2, 16)
    add_footer(slide, len(prs.slides), PROJECT_SHORT)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_slide_bg(slide, "technical")
    add_title_text(slide, "Безопасность и устойчивость")
    add_card(slide, 0.9, 1.9, 2.8, 1.15, "Безопасность", "AES-GCM для payload, RSA-OAEP для ключа, подпись метаданных", fill=PPT_LIGHT_BLUE)
    add_card(slide, 4.0, 1.9, 2.8, 1.15, "Доверие", "Pairing, trusted state, block list, peer identity", fill=PPT_LIGHT_BLUE)
    add_card(slide, 7.1, 1.9, 2.8, 1.15, "Доставка", "Dedup, ACK, retry, route health, relay/failover", fill=PPT_LIGHT_BLUE)
    add_card(slide, 10.2, 1.9, 2.2, 1.15, "Диагностика", "Event log, metrics, topology snapshot", fill=PPT_LIGHT_BLUE)
    add_picture(slide, assets["mobile_profile"], 1.2, 3.5, 2.1)
    add_picture(slide, assets["mobile_chats"], 3.75, 3.5, 2.1)
    add_picture(slide, assets["mobile_transfers"], 6.3, 3.5, 2.1)
    add_picture(slide, assets["desktop_profile"], 8.75, 3.65, 3.2)
    add_footer(slide, len(prs.slides), PROJECT_SHORT)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_slide_bg(slide, "technical")
    add_title_text(slide, "Реализованный прототип")
    add_picture(slide, assets["proto_main"], 0.8, 1.75, 5.7)
    add_picture(slide, assets["proto_inventory_table"], 6.7, 1.75, 5.6)
    add_picture(slide, assets["proto_inventory_edit"], 0.8, 4.55, 5.7)
    add_picture(slide, assets["mobile_home"], 8.9, 4.45, 1.7)
    add_picture(slide, assets["mobile_call"], 10.85, 4.45, 1.55)
    add_footer(slide, len(prs.slides), PROJECT_SHORT)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_slide_bg(slide, "technical")
    add_title_text(slide, "Статус реализации, тесты и ограничения")
    add_card(slide, 0.9, 1.95, 3.25, 1.15, "Статус", "Функциональный прототип", fill=PPT_LIGHT_BLUE)
    add_card(slide, 0.9, 3.35, 3.25, 1.15, "Подтверждено", "MeshNode API, inventory screens, hybrid sync, tests", fill=PPT_LIGHT_BLUE)
    add_card(slide, 0.9, 4.75, 3.25, 1.15, "Ограничения", "Не заявляется внедрение и не завышается готовность до серийного продукта", fill=PPT_LIGHT_BLUE)
    add_body_bullets(slide, [
        "24.03.2026 успешно выполнены CentralHybridServicesTest, InventoryCodeServiceTest и InventoryLabelServiceTest;",
        "central-контур в текущем репозитории подтверждён со стороны клиентской интеграции и sync-логики;",
        "следующий технически корректный шаг - пилотная апробация на ограниченной площадке;",
    ], 4.7, 2.0, 7.0, 3.8, 17)
    add_footer(slide, len(prs.slides), PROJECT_SHORT)

    prs.save(FINAL_FILES["technical_pptx"])


def build_all() -> None:
    ensure_dirs()
    diagrams = build_custom_diagrams()
    extracted = render_existing_slide_images()
    copied = copy_static_images()
    assets = {**diagrams, **extracted, **copied}

    build_doc_from_md(
        "02-osnovnaya-zayavka.md",
        FINAL_FILES["application"],
        "Основная заявка",
        [PROJECT_TITLE, UNIVERSITY, PARTICIPANT],
        title_page=True,
        manual_contents=False,
    )

    note_injections = {
        "Архитектурный подход": [(assets["architecture_existing"], "Рисунок 1. Компонентная архитектура проекта. Источник: техническая презентация проекта.", 16.8)],
        "6.3 Центральный контур": [(assets["central_contour_services"], "Рисунок 1.1. Схема сервисов и взаимодействий в центральном контуре. Источник: подготовлено автором на основе архитектурных материалов проекта и expert-link-core.", 16.8)],
        "Инвентаризация и сессии": [(assets["inventory_er_existing"], "Рисунок 2. Верхнеуровневая ER-модель инвентаризационного контура. Источник: техническая презентация проекта.", 16.8)],
        "Коммуникации и распределённая работа": [(assets["comm_er_existing"], "Рисунок 3. Верхнеуровневая ER-модель коммуникационного контура и синхронизации. Источник: техническая презентация проекта.", 16.8)],
        "Транспорт пакетов": [(assets["packet_sequence_existing"], "Рисунок 4. Последовательность обработки входящего сетевого пакета. Источник: техническая презентация проекта.", 16.8)],
        "Передача файлов": [(assets["file_flow_existing"], "Рисунок 5. Поток передачи файла в распределённом контуре. Источник: техническая презентация проекта.", 16.8)],
        "Сигнальный обмен и медиаканал звонка": [(assets["call_flow_existing"], "Рисунок 6. Поток звонка: сигнальный обмен и медиаканал. Источник: техническая презентация проекта.", 16.8)],
        "Сочетание автономной работы и последующей синхронизации": [(assets["sync"], "Рисунок 7. Локальная очередь изменений и последующая синхронизация. Источник: подготовлено автором на основе архитектурных материалов проекта.", 16.8)],
        "Как система работает в разных режимах связности": [(assets["online_offline"], "Рисунок 8. Режимы работы при наличии и отсутствии внешней связи. Источник: подготовлено автором на основе архитектурных материалов проекта.", 16.8)],
        "Сценарии применения": [(assets["commission"], "Рисунок 9. Сценарий работы комиссии на удалённом объекте. Источник: подготовлено автором на основе проектных материалов.", 16.8)],
    }
    build_doc_from_md(
        "03-poyasnitelnaya-zapiska.md",
        FINAL_FILES["note"],
        "Пояснительная записка",
        [PROJECT_TITLE, UNIVERSITY, PARTICIPANT],
        title_page=True,
        manual_contents=True,
        heading_injections=note_injections,
        chapter_page_breaks=True,
    )

    build_lean_canvas_doc()
    build_summary_doc()
    build_cover_letter_doc()
    build_appendices_doc(assets)
    build_sources_doc()
    build_checklist_doc()
    build_business_pptx(assets)
    build_technical_pptx(assets)
    copy_delivery_aliases()


if __name__ == "__main__":
    build_all()
