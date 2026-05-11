from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.shared import Pt
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR, MSO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt as PptPt


ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = ROOT / "final-package"
ASSET_DIR = OUTPUT_DIR / "_assets"

BASE_PPTX = Path("/home/itech/Downloads/Распределённый_контур_инвентаризации_оптимизировано_1(1).pptx")

OUTPUT_PPTX = OUTPUT_DIR / "Газпром_Презентация_для_защиты_Распределённая_платформа_инвентаризации.pptx"
OUTPUT_DOCX = OUTPUT_DIR / "Газпром_Текст_к_слайдам_Распределённая_платформа_инвентаризации.docx"

ARCHITECTURE_IMAGE = ASSET_DIR / "architecture_existing.png"
PROTO_MAIN_IMAGE = ASSET_DIR / "proto_main.png"
PROTO_INVENTORY_IMAGE = ASSET_DIR / "proto_inventory_cards.png"
SCENARIO_IMAGE = ASSET_DIR / "proto_inventory_edit.png"

BLUE = RGBColor(13, 93, 184)
TEAL = RGBColor(11, 122, 111)
TEXT = RGBColor(39, 50, 61)
GRAY = RGBColor(120, 132, 146)
LIGHT_BLUE = RGBColor(219, 233, 250)
LIGHT_TEAL = RGBColor(228, 244, 241)
LIGHT_GRAY = RGBColor(236, 240, 245)
CARD_LINE = RGBColor(194, 204, 217)
GREEN = RGBColor(191, 222, 167)
WHITE = RGBColor(255, 255, 255)

SLIDE_NOTES = [
    {
        "title": "Титульный слайд",
        "text": (
            "На защите я представляю проект «Распределённая платформа инвентаризации» для удалённых объектов "
            "при нестабильной связи. Это прикладная распределённая информационная система, которая объединяет "
            "учёт имущества, работу комиссии и оперативное взаимодействие сотрудников в одном цифровом контуре.\n\n"
            "Сразу зафиксирую позиционирование проекта: это не отдельный мессенджер и не обычная "
            "централизованная учётная система. Решение ориентировано на удалённые площадки, где нельзя "
            "останавливать процесс из-за деградации внешнего канала."
        ),
    },
    {
        "title": "Общая информация о проекте",
        "text": (
            "Проект объединяет предметный inventory-контур, маркировку, локальную работу, коммуникации и "
            "последующую синхронизацию с центральной системой. То есть сотрудник на площадке получает "
            "не набор разрозненных инструментов, а единый рабочий контур.\n\n"
            "По текущему состоянию это функциональный прототип. Подтверждены desktop- и Android-клиенты, "
            "локальный mesh-контур, inventory-сценарии, интеграция с центральным контуром и профильные тесты, "
            "поэтому проект корректно позиционировать как решение, подготовленное к пилотной апробации."
        ),
    },
    {
        "title": "Ограничения среды и сценарии использования",
        "text": (
            "На удалённых площадках у процесса есть четыре базовых ограничения: территориальная распределённость, "
            "нестабильная или отсутствующая связь с центром, необходимость продолжать учёт локально и "
            "необходимость последующей консолидации после восстановления связи.\n\n"
            "В существующей практике учёт, промежуточные результаты, документы и обсуждения часто "
            "распределены между центральной системой, таблицами, бумажными формами и сторонними "
            "мессенджерами. Именно этот операционный разрыв и является исходной проблемой проекта."
        ),
    },
    {
        "title": "Обзор решения",
        "text": (
            "Проект предлагает единый рабочий контур, в котором инвентаризация и взаимодействие сотрудников "
            "не разделены на разные продукты. Внутри платформы находятся карточки имущества, сессии и проверки, "
            "маркировка, чаты, файлы, обсуждения и последующая синхронизация.\n\n"
            "Ключевая идея здесь в том, что спорный объект можно сразу обсудить, приложить материалы и "
            "зафиксировать решение комиссии, не теряя связь между самим объектом и коммуникацией вокруг него."
        ),
    },
    {
        "title": "Как система работает с интернетом и без него",
        "text": (
            "При наличии связи система синхронизирует данные с центральным контуром и получает обновления. "
            "При отсутствии внешнего канала пользователь продолжает работу локально: сохраняет доступ к "
            "карточкам, сессиям, чатам, файлам и обсуждениям.\n\n"
            "То есть offline здесь не самоцель, а способ не останавливать реальную работу комиссии и сотрудников "
            "на объекте до восстановления внешнего соединения."
        ),
    },
    {
        "title": "Инвентаризационный сценарий на удалённом объекте",
        "text": (
            "Типовой сценарий комиссии начинается с подготовки сессии и списка объектов, затем идёт выезд на "
            "объект, сканирование, открытие карточки, фиксация состояния и расхождений, прикрепление фото "
            "и документов, а затем локальное сохранение результатов с последующей синхронизацией.\n\n"
            "Здесь важно, что замечания, фото и комментарии привязываются к объекту сразу на месте, а "
            "обсуждение спорных случаев ведётся в том же рабочем контуре."
        ),
    },
    {
        "title": "Практическая ценность, новизна и области применения",
        "text": (
            "Практическая ценность проекта в непрерывности работы, снижении числа ручных переносов и в том, "
            "что материалы проверки не теряются вне карточки объекта. Новизна проекта не в отдельных модулях, "
            "а в сочетании автономной локальной работы, встроенных коммуникаций и поздней синхронизации.\n\n"
            "Применение проекта наиболее естественно на удалённых производственных площадках, складах, базах, "
            "в распределённых организациях и в контролируемых пилотах в вузовской среде."
        ),
    },
    {
        "title": "Компонентная архитектура проекта",
        "text": (
            "Архитектура гибридная. Локальный edge-узел сочетает интерфейс, inventory-логику, локальное хранилище, "
            "mesh-коммуникации и очередь синхронизации. Центральный контур используется как канонический источник "
            "данных и точка консолидации.\n\n"
            "Это позволяет не блокировать локальные операции недоступностью центральной системы и при этом "
            "сохранять управляемую интеграцию с gateway, access, inventory, sync, relay, attachment и export сервисами."
        ),
    },
    {
        "title": "Как синхронизируются изменения",
        "text": (
            "Локальная операция сначала фиксируется в локальном контуре и попадает в очередь изменений. После "
            "восстановления связи узел отправляет пакет изменений в central sync с идемпотентным changeId. "
            "Центральный контур подтверждает запись, возвращает новые дельты и при необходимости фиксирует конфликт.\n\n"
            "Важно, что конфликты не теряются молча: они подлежат явному разрешению. За счёт этого поздняя "
            "синхронизация не разрушает локальную работу и не превращает её в ручной перенос."
        ),
    },
    {
        "title": "Последовательность обработки входящего сетевого пакета",
        "text": (
            "Пакет попадает в транспортный слой, затем в NodeLifecycleService, где проходят дедупликация, "
            "проверка доверия и подписи. Только после этого payload передаётся в прикладной сервис.\n\n"
            "Такой подход отделяет сетевую доставку от предметной обработки и централизует критические проверки "
            "устойчивости и безопасности в одной точке жизненного цикла пакета."
        ),
    },
    {
        "title": "Поток передачи файла",
        "text": (
            "Передача файла организована как управляемый протокол: отправитель предлагает файл, получатель "
            "подтверждает приём, затем файл передаётся чанками с подтверждением каждого блока. После сборки "
            "получатель проверяет итоговый SHA-256.\n\n"
            "Если часть чанков потерялась, используется дозагрузка только недостающих фрагментов. Именно это "
            "делает файловый обмен устойчивым при нестабильном соединении."
        ),
    },
    {
        "title": "Устойчивость и безопасность",
        "text": (
            "Устойчивость проекта обеспечивается автономной работой, дедупликацией, повторными попытками, "
            "резервными маршрутами и контролем состояния сети. Для взаимодействия узлов используется trust-модель "
            "через pairing invite, peerId и проверку подписи пакета.\n\n"
            "Полезная нагрузка шифруется через AES-GCM, ключ защищается RSA-OAEP, а для звонков signaling "
            "разделён с медиапотоком, который передаётся по WebRTC с DTLS-SRTP."
        ),
    },
    {
        "title": "Текущий статус и план пилота",
        "text": (
            "Корректная стадия проекта — функциональный прототип. Это означает, что у решения уже есть "
            "экранные сценарии, API, локальный и центральный контуры, механизмы синхронизации и подтверждённый "
            "технический задел, но промышленное внедрение ещё не заявляется как состоявшийся факт.\n\n"
            "Реалистичный следующий шаг — поэтапный пилот: одна площадка, затем offline-сценарии, далее "
            "расширение числа пользователей и подготовка интеграционных требований."
        ),
    },
    {
        "title": "Прототип: рабочий центр",
        "text": (
            "На этом слайде показан рабочий центр прототипа. Видно, что платформа уже поддерживает "
            "оперативную картину по организации, объектам, проверкам и инцидентам, а также быстрый вход "
            "к основным действиям.\n\n"
            "Этот слайд нужен для демонстрации того, что у проекта уже есть прикладной пользовательский слой, "
            "а не только архитектурное описание."
        ),
    },
    {
        "title": "Прототип: inventory-контур",
        "text": (
            "Здесь показан экран inventory-контура со списком оборудования и действиями по карточке, "
            "правке и кодам. Это подтверждает реализованность предметной части проекта и готовность "
            "демонстрировать реальные рабочие сценарии комиссии.\n\n"
            "На практике именно такие экраны обеспечивают быстрый переход от проверки объекта к фиксации "
            "результата и работе с маркировкой."
        ),
    },
    {
        "title": "Прототип: сопряжение и коммуникации",
        "text": (
            "Последний слайд показывает вход в коммуникационный контур: создание приглашения, сопряжение, "
            "QR-сценарий и переход к дальнейшему взаимодействию. Здесь важно подчеркнуть, что коммуникации "
            "являются частью самой платформы, а не внешней интеграцией.\n\n"
            "Итоговый вывод таков: проект уже сейчас существует как инженерно оформленный прототип, который "
            "совмещает прикладную бизнесовую ценность и техническую реализуемость."
        ),
    },
]


def remove_shape(shape) -> None:
    shape._element.getparent().remove(shape._element)


def clear_slide_except(slide, keep_indices: list[int]) -> None:
    keep_shapes = [slide.shapes[index] for index in keep_indices]
    for shape in list(slide.shapes):
        if shape not in keep_shapes:
            remove_shape(shape)


def set_text(shape, text: str) -> None:
    text_frame = shape.text_frame
    text_frame.clear()
    text_frame.word_wrap = True
    for index, paragraph_text in enumerate(text.split("\n")):
        paragraph = text_frame.paragraphs[0] if index == 0 else text_frame.add_paragraph()
        paragraph.text = paragraph_text


def style_text(shape, size: int, color: RGBColor = TEXT, bold: bool = False, align: PP_ALIGN = PP_ALIGN.LEFT) -> None:
    shape.text_frame.word_wrap = True
    for paragraph in shape.text_frame.paragraphs:
        paragraph.alignment = align
        for run in paragraph.runs:
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.color.rgb = color


def set_title(shape, text: str, size: int = 24) -> None:
    set_text(shape, text)
    style_text(shape, size=size, color=BLUE, bold=True)


def add_textbox(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    text: str,
    *,
    size: int = 18,
    color: RGBColor = TEXT,
    bold: bool = False,
    align: PP_ALIGN = PP_ALIGN.LEFT,
    valign: MSO_ANCHOR = MSO_ANCHOR.TOP,
) :
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    shape.text_frame.margin_left = PptPt(4)
    shape.text_frame.margin_right = PptPt(4)
    shape.text_frame.margin_top = PptPt(2)
    shape.text_frame.margin_bottom = PptPt(2)
    shape.text_frame.vertical_anchor = valign
    set_text(shape, text)
    style_text(shape, size=size, color=color, bold=bold, align=align)
    return shape


def add_box(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    text: str,
    *,
    fill: RGBColor = WHITE,
    line: RGBColor = CARD_LINE,
    text_color: RGBColor = TEXT,
    size: int = 16,
    bold: bool = False,
    align: PP_ALIGN = PP_ALIGN.CENTER,
    rounded: bool = True,
    line_width: float = 1.25,
) :
    shape_type = MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if rounded else MSO_AUTO_SHAPE_TYPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    shape.line.width = PptPt(line_width)
    shape.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    shape.text_frame.margin_left = PptPt(6)
    shape.text_frame.margin_right = PptPt(6)
    shape.text_frame.margin_top = PptPt(2)
    shape.text_frame.margin_bottom = PptPt(2)
    set_text(shape, text)
    style_text(shape, size=size, color=text_color, bold=bold, align=align)
    return shape


def add_frame(slide, x: float, y: float, w: float, h: float, line: RGBColor = CARD_LINE) :
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.background()
    shape.line.color.rgb = line
    shape.line.width = PptPt(1.0)
    return shape


def add_line(slide, x1: float, y1: float, x2: float, y2: float, color: RGBColor = GRAY, width: float = 1.5):
    connector = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    connector.line.color.rgb = color
    connector.line.width = PptPt(width)
    return connector


def add_arrow(slide, x1: float, y1: float, x2: float, y2: float, color: RGBColor = TEAL, width: float = 1.5):
    add_line(slide, x1, y1, x2, y2, color=color, width=width)
    size = 0.12
    if abs(x2 - x1) >= abs(y2 - y1):
        if x2 >= x1:
            triangle = slide.shapes.add_shape(
                MSO_AUTO_SHAPE_TYPE.ISOSCELES_TRIANGLE,
                Inches(x2 - size),
                Inches(y2 - size / 2),
                Inches(size),
                Inches(size),
            )
            triangle.rotation = 90
        else:
            triangle = slide.shapes.add_shape(
                MSO_AUTO_SHAPE_TYPE.ISOSCELES_TRIANGLE,
                Inches(x2),
                Inches(y2 - size / 2),
                Inches(size),
                Inches(size),
            )
            triangle.rotation = 270
    else:
        if y2 >= y1:
            triangle = slide.shapes.add_shape(
                MSO_AUTO_SHAPE_TYPE.ISOSCELES_TRIANGLE,
                Inches(x2 - size / 2),
                Inches(y2 - size),
                Inches(size),
                Inches(size),
            )
            triangle.rotation = 180
        else:
            triangle = slide.shapes.add_shape(
                MSO_AUTO_SHAPE_TYPE.ISOSCELES_TRIANGLE,
                Inches(x2 - size / 2),
                Inches(y2),
                Inches(size),
                Inches(size),
            )
            triangle.rotation = 0
    triangle.fill.solid()
    triangle.fill.fore_color.rgb = color
    triangle.line.color.rgb = color
    return triangle


def replace_main_picture(slide, image_path: Path) -> None:
    pictures = [shape for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE]
    main_picture = max(pictures, key=lambda shape: shape.width * shape.height)
    left, top, width, height = main_picture.left, main_picture.top, main_picture.width, main_picture.height
    remove_shape(main_picture)
    slide.shapes.add_picture(str(image_path), left, top, width=width, height=height)


def prepare_content_slide(slide, title: str) -> None:
    clear_slide_except(slide, [0, len(slide.shapes) - 1])
    set_title(slide.shapes[0], title)


def prepare_proto_diagram_slide(slide, title: str) -> None:
    clear_slide_except(slide, [0, 1, 2, 3, 4])
    set_title(slide.shapes[0], title, size=24)


def prepare_proto_image_slide(slide, title: str, caption: str, image_path: Path | None = None) -> None:
    set_title(slide.shapes[0], title, size=24)
    set_text(slide.shapes[5], caption)
    style_text(slide.shapes[5], size=14, color=GRAY, bold=False)
    if image_path is not None:
        replace_main_picture(slide, image_path)


def build_title_slide(slide) -> None:
    set_text(slide.shapes[0], "Титарь Игорь Андреевич")
    set_text(slide.shapes[1], "Защита проекта для олимпиады «Газпром»")
    set_text(slide.shapes[2], "НИЯУ МИФИ, ИИКС, кафедра №22")
    set_text(slide.shapes[3], "Распределённая платформа\nинвентаризации")
    set_text(slide.shapes[4], "15 апреля 2026 г.")
    set_text(slide.shapes[5], "Автор проекта | tigor7750@gmail.com | @ITehn")
    style_text(slide.shapes[0], size=24, color=GRAY)
    style_text(slide.shapes[1], size=22, color=BLUE)
    style_text(slide.shapes[2], size=16, color=TEXT)
    style_text(slide.shapes[3], size=36, color=BLUE)
    style_text(slide.shapes[4], size=15, color=TEXT)
    style_text(slide.shapes[5], size=16, color=TEXT)


def build_overview_slide(slide) -> None:
    set_title(slide.shapes[0], "Общая информация о проекте")
    set_text(
        slide.shapes[1],
        "Проект представляет собой распределённую платформу инвентаризации, обмена данными и оперативного "
        "взаимодействия для удалённых площадок. Система объединяет карточки имущества, инвентаризационные "
        "проверки, маркировку, локальный mesh-контур и встроенные коммуникации. Цель проекта — сохранить "
        "работоспособность учёта и взаимодействия сотрудников при устойчивой связи, её деградации и полном "
        "отсутствии централизованного интернета.",
    )
    set_text(
        slide.shapes[2],
        "Текущая стадия — функциональный прототип. Подтверждены desktop- и Android-клиенты, inventory-контур, "
        "поздняя синхронизация, API и профильные тесты сервисов.",
    )
    set_text(slide.shapes[3], "Что представляет собой решение")
    set_text(slide.shapes[4], "Что уже подтверждено")
    set_text(slide.shapes[5], "Концепт")
    set_text(slide.shapes[6], "Отрасл.\nпилот")
    set_text(slide.shapes[7], "Локал.\nпилот")
    set_text(slide.shapes[8], "Внедрение")
    set_text(slide.shapes[9], "Идея")
    set_text(slide.shapes[10], "Прототип")


def build_constraints_slide(slide) -> None:
    prepare_content_slide(slide, "Ограничения среды и сценарии использования")

    add_box(
        slide,
        0.48,
        1.65,
        6.55,
        4.75,
        "удалённые объекты и территориально\nраспределённые площадки;\n\n"
        "нестабильная или отсутствующая связь с\nцентральным контуром;\n\n"
        "необходимость продолжать учёт и\nвзаимодействие без внешнего канала;\n\n"
        "последующая консолидация данных после\nвосстановления связи;",
        fill=WHITE,
        line=GRAY,
        text_color=TEXT,
        size=18,
        bold=False,
        align=PP_ALIGN.LEFT,
        rounded=False,
        line_width=1.0,
    )

    add_textbox(slide, 7.46, 1.58, 4.15, 0.32, "Разрозненный процесс", size=17, color=BLUE, bold=True)
    add_textbox(
        slide,
        7.46,
        1.86,
        4.2,
        0.22,
        "Учёт, обсуждения и материалы живут в разных инструментах",
        size=9,
        color=GRAY,
    )

    central = add_box(slide, 7.72, 2.3, 1.15, 0.48, "Центральная\nучётная\nсистема", line=BLUE, text_color=BLUE, size=11, bold=True)
    excel = add_box(slide, 9.0, 2.05, 1.45, 0.42, "Excel и\nлокальные\nтаблицы", size=10, bold=True)
    paper = add_box(slide, 9.0, 2.55, 1.45, 0.42, "Бумажные\nформы и\nакты", size=10, bold=True)
    messenger = add_box(slide, 9.0, 3.05, 1.45, 0.42, "Сторонний\nмессенджер", size=10, bold=True)
    commission = add_box(
        slide,
        10.75,
        2.3,
        1.85,
        0.68,
        "Комиссия на\nудалённой\nплощадке",
        fill=LIGHT_TEAL,
        line=TEAL,
        text_color=TEAL,
        size=11,
        bold=True,
    )
    add_arrow(slide, 8.87, 2.55, 9.0, 2.25, color=BLUE)
    add_arrow(slide, 8.87, 2.55, 9.0, 2.75, color=BLUE)
    add_arrow(slide, 8.87, 2.55, 9.0, 3.25, color=BLUE)
    add_arrow(slide, 10.45, 2.25, 10.75, 2.62, color=GRAY)
    add_arrow(slide, 10.45, 2.75, 10.75, 2.64, color=GRAY)
    add_arrow(slide, 10.45, 3.25, 10.75, 2.66, color=GRAY)
    add_textbox(
        slide,
        9.55,
        3.32,
        2.65,
        0.8,
        "Итог: данные, материалы и\nобсуждения расходятся, а при\nпотере связи часть процесса\nуходит в ручной режим.",
        size=10,
        color=TEXT,
    )


def build_solution_slide(slide) -> None:
    prepare_content_slide(slide, "Обзор решения")

    add_frame(slide, 0.32, 1.32, 7.88, 4.84, line=CARD_LINE)
    add_textbox(slide, 0.62, 1.56, 4.4, 0.3, "Единый рабочий контур проекта", size=18, color=BLUE, bold=True)
    add_textbox(
        slide,
        0.62,
        1.84,
        5.2,
        0.22,
        "Инвентаризация и взаимодействие сотрудников объединены в одной платформе",
        size=9,
        color=GRAY,
    )

    add_box(slide, 0.82, 2.44, 1.55, 0.72, "Локальные\nузлы и\nавтономная\nработа", size=12, bold=True)
    add_box(slide, 0.82, 3.5, 1.55, 0.62, "Маркировка,\nQR и\nштрихкоды", size=12, bold=True)
    add_box(
        slide,
        2.78,
        2.06,
        2.28,
        2.18,
        "",
        fill=LIGHT_BLUE,
        line=BLUE,
        text_color=BLUE,
        size=17,
        bold=True,
    )
    add_textbox(
        slide,
        2.96,
        2.28,
        1.92,
        0.48,
        "Распределённая\nплатформа",
        size=17,
        color=BLUE,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_box(slide, 3.14, 3.18, 1.56, 0.44, "Карточки имущества", size=10, bold=True)
    add_box(slide, 3.14, 3.74, 1.56, 0.44, "Сессии и проверки", size=10, bold=True)
    add_box(slide, 5.88, 2.44, 1.55, 0.72, "Чаты,\nфайлы,\nобсуждения\nи звонки", size=12, bold=True)
    add_box(slide, 5.88, 3.5, 1.55, 0.62, "Central sync\nи обмен\nпри связи", size=12, bold=True)
    add_arrow(slide, 2.37, 2.8, 2.78, 2.8, color=TEAL)
    add_arrow(slide, 2.37, 3.81, 2.78, 3.81, color=TEAL)
    add_arrow(slide, 5.06, 2.8, 5.88, 2.8, color=TEAL)
    add_arrow(slide, 5.06, 3.81, 5.88, 3.81, color=TEAL)

    add_textbox(
        slide,
        8.52,
        1.85,
        3.4,
        3.1,
        "edge-узел объединяет inventory,\nлокальную БД и\nmesh-коммуникации;\n\n"
        "central-контур нужен для\nконсолидации и работы с\nканоническими данными;\n\n"
        "локальные операции не ждут\nдоступности центральной\nсистемы;",
        size=16,
        color=TEXT,
    )
    add_box(slide, 8.55, 5.48, 4.0, 0.82, "", fill=LIGHT_BLUE, line=CARD_LINE, rounded=True, line_width=1.0)
    add_textbox(slide, 8.78, 5.56, 3.54, 0.22, "Ключевая идея", size=16, color=BLUE, bold=True, align=PP_ALIGN.CENTER)
    add_textbox(
        slide,
        8.78,
        5.82,
        3.54,
        0.28,
        "Одна платформа для учёта,\nкоммуникаций и синхронизации.",
        size=12,
        color=TEXT,
        align=PP_ALIGN.CENTER,
    )


def build_online_offline_slide(slide) -> None:
    prepare_content_slide(slide, "Как система работает с интернетом и без него")

    add_frame(slide, 0.36, 1.45, 7.92, 4.85, line=GRAY)
    add_textbox(slide, 0.7, 1.68, 3.8, 0.28, "Два рабочих режима", size=17, color=BLUE, bold=True)
    add_textbox(
        slide,
        0.7,
        1.96,
        5.1,
        0.22,
        "Отсутствие внешнего канала не останавливает учёт, проверку и взаимодействие",
        size=9,
        color=GRAY,
    )

    add_box(slide, 0.82, 2.15, 2.95, 2.52, "", fill=LIGHT_BLUE, line=BLUE, text_color=BLUE, size=16, bold=True)
    add_box(slide, 4.54, 2.15, 2.95, 2.52, "", fill=LIGHT_TEAL, line=TEAL, text_color=TEAL, size=16, bold=True)
    add_textbox(slide, 1.2, 2.32, 2.2, 0.24, "Online-режим", size=16, color=BLUE, bold=True, align=PP_ALIGN.CENTER)
    add_textbox(slide, 4.96, 2.32, 2.1, 0.24, "Offline-режим", size=16, color=TEAL, bold=True, align=PP_ALIGN.CENTER)

    online_items = [
        "локальная работа пользователя",
        "доступ к карточкам и сессиям",
        "обсуждения, файлы и звонки",
        "синхронизация с центральным\nконтуром",
        "консолидация и обновление данных",
    ]
    offline_items = [
        "локальная автономная работа",
        "взаимодействие между соседними\nузлами",
        "фиксация результатов и материалов\nна месте",
        "накопление изменений в очереди",
        "поздняя синхронизация после\nвосстановления связи",
    ]

    for index, item in enumerate(online_items):
        add_box(slide, 1.05, 2.76 + index * 0.35, 2.48, 0.26, item, size=9, bold=True, line=CARD_LINE)
    for index, item in enumerate(offline_items):
        add_box(slide, 4.77, 2.76 + index * 0.35, 2.48, 0.26, item, size=9, bold=True, line=CARD_LINE)

    add_textbox(
        slide,
        8.55,
        2.05,
        3.35,
        3.2,
        "при наличии связи изменения\nуходят в central sync;\n\n"
        "при потере внешнего канала\nкомиссия продолжает работу\nлокально;\n\n"
        "после восстановления связи\nнакопленные данные\nконсолидируются автоматически;",
        size=16,
        color=TEXT,
    )


def build_scenario_slide(slide) -> None:
    prepare_content_slide(slide, "Инвентаризационный сценарий на удалённом\nобъекте")

    add_frame(slide, 0.42, 1.76, 7.62, 4.56, line=CARD_LINE)
    add_textbox(slide, 0.72, 1.98, 3.6, 0.28, "Сценарий работы комиссии", size=18, color=BLUE, bold=True)
    add_textbox(slide, 0.72, 2.24, 4.2, 0.2, "От подготовки сессии до последующей синхронизации", size=9, color=GRAY)

    steps = [
        ("1. Подготовка сессии", LIGHT_BLUE, BLUE, BLUE),
        ("2. Выезд на объект", WHITE, CARD_LINE, TEXT),
        ("3. Сканирование и открытие карточки", WHITE, CARD_LINE, TEXT),
        ("4. Проверка состояния и фиксация расхождений", WHITE, CARD_LINE, TEXT),
        ("5. Фото, документы и комментарии", WHITE, CARD_LINE, TEXT),
        ("6. Локальное сохранение и sync", WHITE, BLUE, BLUE),
    ]
    y = 2.62
    for index, (label, fill_color, line_color, text_color) in enumerate(steps):
        add_box(slide, 0.92, y, 5.96, 0.42, label, fill=fill_color, line=line_color, text_color=text_color, size=13, bold=True, align=PP_ALIGN.LEFT)
        if index < len(steps) - 1:
            add_arrow(slide, 3.9, y + 0.42, 3.9, y + 0.54, color=TEAL)
        y += 0.56

    add_textbox(
        slide,
        8.32,
        1.94,
        3.55,
        2.6,
        "сканирование объекта сразу\nоткрывает карточку и действия\nпо проверке;\n\n"
        "замечания, фото и\nкомментарии привязываются к\nобъекту на месте;\n\n"
        "спорные случаи обсуждаются\nв том же рабочем контуре;",
        size=15,
        color=TEXT,
    )
    slide.shapes.add_picture(str(SCENARIO_IMAGE), Inches(9.02), Inches(4.86), width=Inches(3.02))


def build_value_slide(slide) -> None:
    prepare_content_slide(slide, "Практическая ценность, новизна и области\nприменения")

    cards = [
        (
            0.62,
            "Практическая\nценность",
            "непрерывность работы\nдаже при потере связи;\n\n"
            "меньше ручного переноса\nмежду инструментами;\n\n"
            "материалы проверки и\nрешения не теряются\nвне карточки объекта;",
        ),
        (
            4.18,
            "Новизна\nрешения",
            "автономная локальная\nработа как базовый\nсценарий;\n\n"
            "встроенные коммуникации\nвнутри предметного\nпроцесса;\n\n"
            "поздняя синхронизация\nвместо ручного переноса\nданных;",
        ),
        (
            7.74,
            "Где применять",
            "удалённые базы,\nсклады и площадки;\n\n"
            "распределённые\nорганизации и службы;\n\n"
            "контролируемый пилот\nв вузе или лаборатории;",
        ),
    ]

    for x, title, body in cards:
        add_box(slide, x, 1.82, 3.18, 3.42, "", fill=WHITE, line=CARD_LINE, rounded=True)
        add_textbox(slide, x + 0.18, 2.05, 2.82, 0.52, title, size=18, color=BLUE, bold=True, align=PP_ALIGN.CENTER)
        add_textbox(slide, x + 0.18, 2.64, 2.82, 2.24, body, size=15, color=TEXT)

    ribbon = add_box(
        slide,
        0.85,
        5.58,
        10.98,
        0.7,
        "Для задач Газпром проект релевантен как распределённая прикладная система для учёта, "
        "локальной устойчивости и работы с данными непосредственно на месте.",
        fill=LIGHT_BLUE,
        line=CARD_LINE,
        text_color=TEXT,
        size=14,
        rounded=True,
        line_width=1.0,
    )
    ribbon.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE


def build_architecture_slide(slide) -> None:
    prepare_content_slide(slide, "Компонентная архитектура проекта")

    slide.shapes.add_picture(str(ARCHITECTURE_IMAGE), Inches(0.42), Inches(1.46), width=Inches(7.55))

    add_textbox(
        slide,
        8.28,
        1.74,
        3.82,
        0.3,
        "Что важно в архитектуре",
        size=18,
        color=BLUE,
        bold=True,
    )
    add_textbox(
        slide,
        8.28,
        2.14,
        3.72,
        2.95,
        "edge-узел включает UI,\ninventory, локальную БД,\nmesh-модуль и очередь sync;\n\n"
        "central-контур содержит gateway,\norganization access, inventory,\nsync, relay, attachment и export;\n\n"
        "локальные операции идут автономно,\nа центральная система остаётся\nточкой консолидации;",
        size=15,
        color=TEXT,
    )
    add_box(
        slide,
        8.36,
        5.68,
        3.52,
        0.54,
        "Гибридная модель = локальная устойчивость + central source of truth",
        fill=LIGHT_BLUE,
        line=CARD_LINE,
        text_color=TEXT,
        size=12,
        bold=True,
        rounded=True,
    )


def build_sync_slide(slide) -> None:
    prepare_content_slide(slide, "Как синхронизируются изменения")

    add_frame(slide, 0.42, 1.5, 7.45, 4.72, line=CARD_LINE)
    add_textbox(slide, 0.72, 1.76, 4.0, 0.28, "Путь изменения до central sync", size=17, color=BLUE, bold=True)
    add_textbox(
        slide,
        0.72,
        2.02,
        4.6,
        0.2,
        "Локальная операция фиксируется сразу, а потом уходит в управляемую синхронизацию",
        size=9,
        color=GRAY,
    )

    boxes = [
        (0.86, 3.04, 1.16, 0.58, "Операция\nна площадке", LIGHT_BLUE, BLUE, BLUE),
        (2.16, 3.04, 1.02, 0.58, "Локальная\nБД", WHITE, CARD_LINE, TEXT),
        (3.34, 3.04, 1.14, 0.58, "Очередь\nизменений", WHITE, CARD_LINE, TEXT),
        (4.64, 3.04, 1.04, 0.58, "Sync\nпакет", WHITE, CARD_LINE, TEXT),
        (5.88, 2.72, 1.42, 0.72, "Central\nSync Service", LIGHT_TEAL, TEAL, TEAL),
        (5.78, 4.02, 1.56, 0.82, "Подтверждение\nи новые\ndelta", LIGHT_BLUE, BLUE, BLUE),
    ]
    for x, y, w, h, label, fill, line, text_color in boxes:
        add_box(slide, x, y, w, h, label, fill=fill, line=line, text_color=text_color, size=11, bold=True)

    add_arrow(slide, 1.96, 3.3, 2.12, 3.3, color=TEAL)
    add_arrow(slide, 3.12, 3.3, 3.28, 3.3, color=TEAL)
    add_arrow(slide, 4.38, 3.3, 4.56, 3.3, color=TEAL)
    add_arrow(slide, 5.61, 3.3, 5.8, 3.03, color=TEAL)
    add_arrow(slide, 6.45, 3.4, 6.45, 4.06, color=TEAL)
    add_arrow(slide, 5.78, 4.44, 3.98, 4.44, color=BLUE)
    add_arrow(slide, 3.98, 4.44, 3.98, 3.66, color=BLUE)
    add_textbox(slide, 4.18, 4.14, 1.25, 0.18, "cursor / delta", size=8, color=GRAY, align=PP_ALIGN.CENTER)
    add_textbox(slide, 5.92, 2.42, 1.32, 0.18, "changeId / idempotency", size=8, color=GRAY, align=PP_ALIGN.CENTER)

    add_textbox(
        slide,
        8.26,
        1.84,
        3.52,
        3.55,
        "операция фиксируется сразу и\nпопадает в локальную очередь;\n\n"
        "узел отправляет sync-пакет с\nидемпотентным changeId;\n\n"
        "центр подтверждает запись,\nвозвращает новые delta и\nвыделяет конфликты отдельно;\n\n"
        "поэтому поздняя sync не\nсводится к ручному переносу;",
        size=15,
        color=TEXT,
    )


def build_packet_slide(slide) -> None:
    prepare_proto_diagram_slide(slide, "Последовательность обработки входящего сетевого пакета")

    add_frame(slide, 0.38, 1.52, 7.98, 4.82, line=CARD_LINE)
    participants = [
        ("Узел A", 0.62),
        ("Transport", 2.02),
        ("Lifecycle", 3.54),
        ("Use-case", 5.06),
        ("Узел B", 6.58),
    ]
    for label, x in participants:
        add_box(slide, x, 1.95, 1.02, 0.38, label, size=11, bold=True)
        add_line(slide, x + 0.51, 2.33, x + 0.51, 5.85, color=CARD_LINE, width=1.0)

    add_arrow(slide, 1.13, 2.82, 2.02, 2.82, color=TEAL)
    add_textbox(slide, 0.86, 2.62, 1.42, 0.18, "POST /api/v1/packets", size=8, color=GRAY, align=PP_ALIGN.CENTER)
    add_arrow(slide, 2.53, 3.18, 3.54, 3.18, color=TEAL)
    add_textbox(slide, 2.42, 2.98, 1.48, 0.18, "handleIncomingPacket", size=8, color=GRAY, align=PP_ALIGN.CENTER)
    add_arrow(slide, 4.05, 3.64, 5.06, 3.64, color=TEAL)
    add_textbox(slide, 4.16, 3.46, 0.98, 0.16, "payload", size=8, color=GRAY, align=PP_ALIGN.CENTER)
    add_arrow(slide, 5.57, 4.02, 4.05, 4.02, color=BLUE)
    add_textbox(slide, 4.47, 3.82, 0.72, 0.18, "результат", size=8, color=GRAY, align=PP_ALIGN.CENTER)
    add_arrow(slide, 3.54, 4.42, 2.53, 4.42, color=BLUE)
    add_textbox(slide, 2.73, 4.22, 0.82, 0.18, "202 / 502", size=8, color=GRAY, align=PP_ALIGN.CENTER)
    add_arrow(slide, 2.53, 4.86, 7.09, 4.86, color=BLUE)
    add_textbox(slide, 3.76, 4.66, 2.16, 0.18, "ответ доставки", size=8, color=GRAY, align=PP_ALIGN.CENTER)

    add_textbox(slide, 3.72, 3.18, 1.02, 0.16, "dedup / trust", size=8, color=GRAY, align=PP_ALIGN.CENTER)
    add_textbox(
        slide,
        8.56,
        1.92,
        3.42,
        4.25,
        "1. Узел A отправляет POST-пакет\nв transport.\n\n"
        "2. Transport передаёт его в\nNodeLifecycleService.\n\n"
        "3. В lifecycle проходят dedup,\ntrust и проверка подписи.\n\n"
        "4. Затем payload уходит в\nприкладной сервис.\n\n"
        "5. Use-case возвращает\nрезультат обработки.\n\n"
        "6. Lifecycle и transport формируют\nответ доставки.",
        size=14,
        color=TEXT,
    )


def build_file_flow_slide(slide) -> None:
    prepare_proto_diagram_slide(slide, "Поток передачи файла")

    add_frame(slide, 0.46, 1.62, 7.62, 4.55, line=CARD_LINE)
    add_box(slide, 1.72, 2.05, 1.02, 0.42, "Узел A", size=12, bold=True)
    add_box(slide, 5.92, 2.05, 1.02, 0.42, "Узел B", size=12, bold=True)
    add_line(slide, 2.23, 2.47, 2.23, 5.74, color=CARD_LINE, width=1.0)
    add_line(slide, 6.43, 2.47, 6.43, 5.74, color=CARD_LINE, width=1.0)

    flows = [
        (2.23, 2.86, 6.43, 2.86, "FILE_OFFER"),
        (6.43, 3.18, 2.23, 3.18, "FILE_ACCEPT"),
        (2.23, 3.72, 6.43, 3.72, "FILE_CHUNK"),
        (6.43, 4.08, 2.23, 4.08, "FILE_ACK"),
        (6.43, 4.82, 2.23, 4.82, "FILE_COMPLETE (SHA-256 ok)"),
        (6.43, 5.18, 2.23, 5.18, "FILE_RESUME_REQUEST"),
    ]
    for x1, y1, x2, y2, label in flows:
        add_arrow(slide, x1, y1, x2, y2, color=TEAL if "FILE_" in label[:10] else BLUE)
        add_textbox(slide, 3.0, y1 - 0.18, 2.6, 0.18, label, size=10, color=TEXT, align=PP_ALIGN.CENTER)

    loop = add_frame(slide, 2.07, 3.42, 4.52, 0.95, line=RGBColor(206, 196, 244))
    loop.line.dash_style = 4
    add_textbox(slide, 2.22, 3.3, 1.0, 0.15, "Цикл по чанкам", size=8, color=GRAY, bold=True)
    add_textbox(slide, 4.02, 3.3, 0.82, 0.15, "[chunks]", size=8, color=GRAY, bold=True, align=PP_ALIGN.CENTER)

    add_textbox(
        slide,
        8.2,
        1.95,
        3.75,
        4.3,
        "1. Отправитель инициирует\nобмен через FILE_OFFER.\n\n"
        "2. Получатель отвечает\nFILE_ACCEPT.\n\n"
        "3. Файл идёт чанками через\nFILE_CHUNK.\n\n"
        "4. Каждый блок подтверждается\nсообщением FILE_ACK.\n\n"
        "5. После сборки проверяется\nитоговый SHA-256.\n\n"
        "6. При пропусках используется\nFILE_RESUME_REQUEST.",
        size=15,
        color=TEXT,
    )


def build_security_slide(slide) -> None:
    prepare_proto_diagram_slide(slide, "Устойчивость и безопасность")

    cards = [
        (0.52, 1.7, "Автономная работа", "локальный узел продолжает работу\nи хранит операционное состояние"),
        (4.25, 1.7, "Доверие узлов", "pairing invite, TTL, nonce,\npeerId и trust state"),
        (7.98, 1.7, "Надёжная доставка", "packetId, dedup, ACK,\nretry и дозагрузка чанков"),
        (0.52, 3.62, "Шифрование данных", "AES-GCM для payload,\nRSA-OAEP для ключа"),
        (4.25, 3.62, "Резервные маршруты", "relay / multihop, route health\nи деградация связности"),
        (7.98, 3.62, "Звонки и media", "CALL_SIGNAL отдельно от media,\nWebRTC + DTLS-SRTP"),
    ]
    for x, y, title, body in cards:
        add_box(slide, x, y, 3.1, 1.45, "", fill=WHITE, line=CARD_LINE, rounded=True)
        add_textbox(slide, x + 0.16, y + 0.12, 2.8, 0.28, title, size=17, color=BLUE, bold=True, align=PP_ALIGN.CENTER)
        add_textbox(slide, x + 0.16, y + 0.52, 2.78, 0.7, body, size=13, color=TEXT, align=PP_ALIGN.CENTER)

    add_box(
        slide,
        1.25,
        5.55,
        10.1,
        0.62,
        "Вместе эти механизмы позволяют не останавливать рабочий процесс при сетевой деградации и "
        "не терять корректность обмена между узлами.",
        fill=LIGHT_BLUE,
        line=CARD_LINE,
        text_color=TEXT,
        size=13,
        bold=True,
        rounded=True,
    )


def build_status_slide(slide) -> None:
    prepare_proto_diagram_slide(slide, "Текущий статус и план пилота")

    add_box(
        slide,
        0.58,
        1.76,
        3.24,
        1.1,
        "Статус проекта\nФункциональный прототип",
        fill=GREEN,
        line=RGBColor(127, 170, 88),
        text_color=TEXT,
        size=18,
        bold=True,
        rounded=True,
    )
    add_box(
        slide,
        4.15,
        1.76,
        4.05,
        1.1,
        "Подтверждено в реализации\nDesktop и Android UI, inventory-контур,\nmesh-коммуникации, hybrid sync, API и тесты",
        fill=LIGHT_BLUE,
        line=CARD_LINE,
        text_color=TEXT,
        size=13,
        bold=False,
        rounded=True,
    )
    add_box(
        slide,
        8.52,
        1.76,
        3.06,
        1.1,
        "Цель следующего этапа\nПилотная апробация на ограниченной\nплощадке без завышения стадии",
        fill=LIGHT_TEAL,
        line=TEAL,
        text_color=TEXT,
        size=13,
        rounded=True,
    )
    add_textbox(slide, 0.62, 3.12, 2.8, 0.28, "Этапы пилота", size=18, color=BLUE, bold=True)

    phases = [
        ("1. Одна площадка", "реестр объектов,\nмаркировка,\n1-2 сессии"),
        ("2. Offline-сценарии", "несколько узлов,\nлокальный обмен,\noffline/online-переходы"),
        ("3. Расширение", "вложения,\nотчёты,\nролевой контекст"),
        ("4. Интеграция", "регламент обмена,\ncentral-контур,\nограничения эксплуатации"),
    ]
    x = 0.64
    for title, body in phases:
        add_box(slide, x, 3.72, 2.76, 1.72, "", fill=WHITE, line=CARD_LINE, rounded=True)
        add_textbox(slide, x + 0.12, 3.88, 2.5, 0.28, title, size=16, color=BLUE, bold=True, align=PP_ALIGN.CENTER)
        add_textbox(slide, x + 0.12, 4.28, 2.48, 0.82, body, size=13, color=TEXT, align=PP_ALIGN.CENTER)
        x += 2.96

    add_box(
        slide,
        1.15,
        5.72,
        10.15,
        0.48,
        "Реалистичная формулировка стадии: не промышленная эксплуатация, а функциональный прототип, готовый к пилотной апробации.",
        fill=WHITE,
        line=GRAY,
        text_color=TEXT,
        size=12,
        rounded=True,
        line_width=1.0,
    )


def build_speaker_doc() -> None:
    document = Document()
    title = document.add_paragraph()
    title_run = title.add_run("Текст выступления по слайдам")
    title_run.bold = True
    title_run.font.size = Pt(16)

    subtitle = document.add_paragraph()
    subtitle_run = subtitle.add_run("Проект: «Распределённая платформа инвентаризации»")
    subtitle_run.font.size = Pt(12)

    intro = document.add_paragraph()
    intro.add_run(
        "Документ подготовлен для устного выступления на защите проекта в рамках Студенческой олимпиады «Газпром» "
        "по профилю «Информационные системы и технологии»."
    )

    for slide_number, slide_data in enumerate(SLIDE_NOTES, start=1):
        document.add_page_break()
        heading = document.add_paragraph()
        heading_run = heading.add_run(f"Слайд {slide_number}. {slide_data['title']}")
        heading_run.bold = True
        heading_run.font.size = Pt(14)

        for block in slide_data["text"].split("\n\n"):
            paragraph = document.add_paragraph(block)
            paragraph.paragraph_format.space_after = Pt(8)

    document.save(str(OUTPUT_DOCX))


def delete_slide(prs: Presentation, index: int) -> None:
    slide_id = prs.slides._sldIdLst[index]
    rel_id = slide_id.rId
    prs.part.drop_rel(rel_id)
    del prs.slides._sldIdLst[index]


def main() -> None:
    build_speaker_doc()

    prs = Presentation(str(BASE_PPTX))

    build_title_slide(prs.slides[0])
    build_overview_slide(prs.slides[1])
    build_constraints_slide(prs.slides[2])
    build_solution_slide(prs.slides[3])
    build_online_offline_slide(prs.slides[4])
    build_scenario_slide(prs.slides[5])
    build_value_slide(prs.slides[6])
    build_architecture_slide(prs.slides[7])
    build_sync_slide(prs.slides[8])
    build_packet_slide(prs.slides[9])
    build_file_flow_slide(prs.slides[10])
    build_security_slide(prs.slides[11])
    build_status_slide(prs.slides[12])

    prepare_proto_image_slide(
        prs.slides[13],
        "Прототип: рабочий центр",
        "Рабочий центр: организация, сводка по объектам, проверки и быстрый вход к основным действиям.",
        PROTO_MAIN_IMAGE,
    )
    prepare_proto_image_slide(
        prs.slides[14],
        "Прототип: inventory-контур",
        "Список оборудования, карточки, операции по правке и работа с кодами в предметном контуре.",
        PROTO_INVENTORY_IMAGE,
    )
    prepare_proto_image_slide(
        prs.slides[19],
        "Прототип: сопряжение и коммуникации",
        "Сопряжение, обмен приглашением и вход в коммуникационный контур внутри той же платформы.",
        None,
    )

    for index in range(len(prs.slides) - 1, -1, -1):
        if index in set(range(0, 15)) | {19}:
            continue
        delete_slide(prs, index)

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUTPUT_PPTX))


if __name__ == "__main__":
    main()
